"""BACKEND-155：app/services/student_import_service.py（preview：解析 Excel、逐列驗證、不寫入；
不受信任 xlsx 的格數 / 解壓防護在交給 openpyxl 之前生效）。
BACKEND-156：execute（重跑 preview、有錯誤列 409 import_has_errors 且不寫入、全部合法以 151 的寫入
邏輯逐列新增（密文 + HMAC）、競態衝突整批回滾 409 import_conflict、稽核 student.import、
不 commit）。"""

from __future__ import annotations

import io
import re
import tracemalloc
import warnings
import zipfile
from collections.abc import Callable, Iterator, Sequence
from datetime import date
from uuid import uuid4

import openpyxl.xml
import pytest
from openpyxl import Workbook
from sqlalchemy import Engine, func, select
from sqlalchemy.orm import Session

from app.api.deps import CurrentStaff
from app.core.config import get_settings
from app.core.crypto import decrypt_bytes, derive_key, encrypt_bytes
from app.core.errors import AppError
from app.core.request_meta import RequestMeta
from app.core.uploads import ValidatedUpload
from app.models.audit import AuditLog
from app.models.students import Student
from app.schemas.students import StudentCreateIn
from app.services import student_import_service
from app.services.student_import_service import (
    IMPORT_COLUMNS,
    MAX_IMPORT_ROWS,
    MAX_XML_ELEMENTS,
    MAX_ZIP_ENTRIES,
    ImportPreview,
    ImportResult,
    ImportRowResult,
    check_header,
    execute,
    preview,
)
from app.services.students.id_number import id_number_hmac
from tests.support.factories import make_class, make_school, make_student
from tests.support.fake_clock import FakeClock

_ID_A = "A123456789"
_ID_B = "B123456708"
_META = RequestMeta(ip="127.0.0.1", user_agent="pytest", request_id="req-1")


@pytest.fixture(autouse=True)
def _crypto_env(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    monkeypatch.setenv("APP_ENV", "test")
    monkeypatch.setenv("DATABASE_URL", "postgresql+psycopg://u:p@127.0.0.1:54342/postgres")
    monkeypatch.setenv("APP_SECRET_KEY", "s" * 48)
    monkeypatch.setenv("PUBLIC_BASE_URL", "http://127.0.0.1:5341")
    monkeypatch.setenv("R2_ENDPOINT_URL", "http://127.0.0.1:54344")
    monkeypatch.setenv("R2_ACCESS_KEY_ID", "afterschool")
    monkeypatch.setenv("R2_SECRET_ACCESS_KEY", "afterschool-local-secret")
    monkeypatch.setenv("R2_BUCKET", "afterschool-local")
    get_settings.cache_clear()
    derive_key.cache_clear()
    yield
    get_settings.cache_clear()
    derive_key.cache_clear()


def _actor(*permissions: str) -> CurrentStaff:
    return CurrentStaff(
        id=uuid4(),
        username="clerk",
        display_name="陳行政",
        role_id=uuid4(),
        role_code="clerk",
        role_name="行政",
        permissions=frozenset(permissions),
        must_change_password=False,
        token_version=0,
    )


_WRITER = _actor("students:write")
_SENSITIVE = _actor("students:write", "students:sensitive")


def _xlsx(rows: Sequence[Sequence[object]], header: Sequence[str] | None = None) -> ValidatedUpload:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    if header is not None:
        ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()
    return ValidatedUpload(
        content=content,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ext="xlsx",
        size=len(content),
    )


def _row(  # 測試資料列：欄位順序對齊 IMPORT_COLUMNS
    student_no: object = "S115101",
    name: object = "林小安",
    gender: object = "男",
    birthday: object = date(2018, 5, 1),
    grade: object = 2,
    school: object = "新生",
    school_class: object = "二年三班",
    klass: object = "低年級 A 班",
    status: object = "在學",
    enrolled: object = "2026-09-01",
    note: object = None,
    id_number: object = None,
    health_note: object = None,
) -> list[object]:
    return [
        student_no,
        name,
        gender,
        birthday,
        grade,
        school,
        school_class,
        klass,
        status,
        enrolled,
        note,
        id_number,
        health_note,
    ]


@pytest.fixture
def lookups(db_session: Session) -> dict[str, object]:
    school = make_school(db_session, name="新生國小")
    school.short_name = "新生"
    inactive = make_school(db_session, name="停用國小")
    inactive.is_active = False
    klass = make_class(db_session, name="低年級 A 班", academic_year=115)
    make_class(db_session, name="低年級 B 班", academic_year=114)  # 別的學年
    make_class(db_session, name="舊班", academic_year=115, archived=True)
    db_session.flush()
    return {"school": school, "class": klass}


def _preview(
    db: Session, upload: ValidatedUpload, actor: CurrentStaff = _SENSITIVE
) -> ImportPreview:
    return preview(db, upload, academic_year=115, actor=actor)


def test_import_preview_valid(db_session: Session, lookups: dict[str, object]) -> None:
    upload = _xlsx(
        [
            _row(id_number=f" {_ID_A.lower()} ", health_note="氣喘"),
            _row(
                student_no="S115102", name="陳小華", gender="女", school="新生國小", status="暫停"
            ),
            _row(
                student_no=115103,  # 數字儲存格
                name="黃小美",
                gender="其他",
                birthday="2017/03/15",
                grade=3.0,
                school=None,
                klass=None,
                status=None,
                enrolled=date(2025, 9, 1),
                note="備註",
            ),
        ],
        header=IMPORT_COLUMNS,
    )

    result = _preview(db_session, upload)

    assert isinstance(result, ImportPreview)
    assert (result.total, result.valid, result.invalid) == (3, 3, 0)
    first, second, third = result.rows
    assert isinstance(first, ImportRowResult)
    assert [r.row_number for r in result.rows] == [2, 3, 4]
    assert all(r.errors == [] for r in result.rows)
    assert isinstance(first.data, StudentCreateIn)
    assert first.data.student_no == "S115101"
    assert first.data.gender == "male"
    assert first.data.birthday == date(2018, 5, 1)
    assert first.data.class_id == lookups["class"].id  # type: ignore[attr-defined]
    assert first.data.school_id == lookups["school"].id  # type: ignore[attr-defined]
    assert first.data.school_class == "二年三班"
    assert first.data.status == "active"
    assert first.data.enrolled_on == date(2026, 9, 1)
    assert first.data.id_number == _ID_A  # 正規化（去空白、大寫）
    assert first.data.health_note == "氣喘"
    assert first.display["學號*"] == "S115101"
    assert first.display["身分證字號"] == "A12****789"
    assert _ID_A not in str(first.display)
    assert "氣喘" not in str(first.display)
    assert second.data is not None
    assert (second.data.gender, second.data.status) == ("female", "suspended")
    assert second.data.school_id == lookups["school"].id  # type: ignore[attr-defined]
    assert third.data is not None
    assert third.data.student_no == "115103"
    assert (third.data.gender, third.data.grade_level) == ("other", 3)
    assert third.data.birthday == date(2017, 3, 15)
    assert third.data.enrolled_on == date(2025, 9, 1)
    assert (third.data.school_id, third.data.class_id, third.data.status) == (None, None, "active")
    assert third.data.note == "備註"
    assert third.display["姓名*"] == "黃小美"


def test_import_preview_header_and_limits(db_session: Session) -> None:
    missing_name = [c for c in IMPORT_COLUMNS if c != "姓名*"]
    with pytest.raises(AppError) as missing:
        _preview(db_session, _xlsx([_row()[:1] + _row()[2:]], header=missing_name))
    assert (missing.value.status, missing.value.code) == (422, "import_invalid_header")
    assert missing.value.details == {"missing": ["姓名*"], "unexpected": []}

    with pytest.raises(AppError) as extra:
        _preview(db_session, _xlsx([[*_row(), 5]], header=[*IMPORT_COLUMNS, "座號"]))
    assert extra.value.code == "import_invalid_header"
    assert extra.value.details == {"missing": [], "unexpected": ["座號"]}
    # 標題去頭尾空白比對；選填欄缺少不算錯
    assert check_header([" 學號* ", "姓名*", "年級*", None, ""]) == ([], [])
    assert check_header(["學號*", "座號", "姓名*"]) == (["年級*"], ["座號"])

    too_many = [_row(student_no=f"S{n:06d}") for n in range(MAX_IMPORT_ROWS + 1)]
    with pytest.raises(AppError) as limit:
        _preview(db_session, _xlsx(too_many, header=IMPORT_COLUMNS))
    assert (limit.value.status, limit.value.code) == (422, "import_too_many_rows")

    with pytest.raises(AppError) as empty:
        _preview(db_session, _xlsx([], header=IMPORT_COLUMNS))
    assert (empty.value.status, empty.value.code) == (422, "import_empty")
    # 全空白列不算資料列；完全沒有標題也是空檔
    with pytest.raises(AppError) as blank:
        _preview(db_session, _xlsx([[None] * len(IMPORT_COLUMNS)], header=IMPORT_COLUMNS))
    assert blank.value.code == "import_empty"
    with pytest.raises(AppError) as no_header:
        _preview(db_session, _xlsx([]))
    assert no_header.value.code == "import_empty"


def test_import_preview_lookup_errors(db_session: Session, lookups: dict[str, object]) -> None:
    upload = _xlsx(
        [
            _row(school="不存在國小"),
            _row(student_no="S115102", klass="不存在班"),
            _row(student_no="S115103", school="停用國小", klass="低年級 B 班"),
            _row(student_no="S115104", klass="舊班"),
            _row(student_no="S115105"),
        ],
        header=IMPORT_COLUMNS,
    )

    result = _preview(db_session, upload)

    assert (result.total, result.valid, result.invalid) == (5, 1, 4)
    assert "找不到國小：不存在國小" in result.rows[0].errors
    assert result.rows[0].data is None
    assert any("找不到安親班班級" in e for e in result.rows[1].errors)
    assert any("找不到國小" in e for e in result.rows[2].errors)  # 停用的國小不算
    assert any("找不到安親班班級" in e for e in result.rows[2].errors)  # 別的學年
    assert any("找不到安親班班級" in e for e in result.rows[3].errors)  # 已封存
    assert result.rows[4].errors == []
    assert result.rows[4].data is not None


def test_import_preview_duplicates(db_session: Session, lookups: dict[str, object]) -> None:
    make_student(db_session, student_no="S115001")
    archived = make_student(db_session, student_no="OLD001", archived=True)
    archived.id_number_enc = encrypt_bytes(_ID_A)
    archived.id_number_hmac = id_number_hmac(_ID_A)
    db_session.flush()
    upload = _xlsx(
        [
            _row(student_no="S115001"),
            _row(student_no="S115200", name="甲"),
            _row(student_no="S115200", name="乙"),
            _row(student_no="S115201", id_number=_ID_A),
            _row(student_no="S115202", id_number=_ID_B),
            _row(student_no="S115203", id_number=_ID_B.lower()),
            _row(student_no="S115204", id_number="A123"),
            _row(student_no="S 115205"),
        ],
        header=IMPORT_COLUMNS,
    )

    result = _preview(db_session, upload)

    rows = result.rows
    assert "學號已存在" in rows[0].errors
    assert "學號在檔案中重複" in rows[1].errors
    assert "學號在檔案中重複" in rows[2].errors
    assert "身分證字號已存在" in rows[3].errors
    assert "身分證字號在檔案中重複" in rows[4].errors
    assert "身分證字號在檔案中重複" in rows[5].errors
    assert any("身分證字號" in e and "格式" in e for e in rows[6].errors)
    assert any("學號" in e for e in rows[7].errors)
    assert result.valid == 0
    assert result.invalid == 8


def test_import_preview_sensitive_permission(
    db_session: Session, lookups: dict[str, object]
) -> None:
    upload = _xlsx(
        [
            _row(id_number=_ID_A),
            _row(student_no="S115102", health_note="對花生過敏"),
            _row(student_no="S115103"),
        ],
        header=IMPORT_COLUMNS,
    )

    result = _preview(db_session, upload, actor=_WRITER)

    assert "沒有權限匯入敏感欄位" in result.rows[0].errors
    assert "沒有權限匯入敏感欄位" in result.rows[1].errors
    assert result.rows[2].errors == []
    assert result.rows[0].display["身分證字號"] == "A12****789"
    assert "對花生過敏" not in str(result.rows[1].display)
    assert result.rows[1].display["健康備註"] != ""
    assert (result.valid, result.invalid) == (1, 2)
    # 有權限時同一檔全數通過
    allowed = _preview(db_session, upload, actor=_SENSITIVE)
    assert allowed.invalid == 0


def test_import_preview_no_writes(db_session: Session, lookups: dict[str, object]) -> None:
    before = db_session.execute(select(func.count()).select_from(Student)).scalar_one()
    upload = _xlsx([_row(), _row(student_no="S115102", id_number=_ID_A)], header=IMPORT_COLUMNS)

    result = _preview(db_session, upload)

    assert result.valid == 2
    assert not db_session.new
    assert db_session.execute(select(func.count()).select_from(Student)).scalar_one() == before


# --- 不受信任 xlsx 的防護（review-r8-b 兩輪打回的攻擊面清單）---------------------------

_SHEET1 = "xl/worksheets/sheet1.xml"
_INVALID = (422, "import_invalid_file")


def _parts(upload: ValidatedUpload) -> dict[str, bytes]:
    src = zipfile.ZipFile(io.BytesIO(upload.content))
    return {i.filename: src.read(i) for i in src.infolist()}


def _pack(parts: dict[str, bytes], *, compress: int = zipfile.ZIP_DEFLATED) -> ValidatedUpload:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compress, compresslevel=9) as dst:
        for name, data in parts.items():
            dst.writestr(name, data)
    content = out.getvalue()
    return ValidatedUpload(
        content=content,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ext="xlsx",
        size=len(content),
    )


def _rewrite_sheet(upload: ValidatedUpload, fn: Callable[[bytes], bytes]) -> ValidatedUpload:
    """改寫第一個工作表的 XML 後重新打包（高壓縮比，模擬小檔案解壓成大量內容）。"""
    parts = _parts(upload)
    parts[_SHEET1] = fn(parts[_SHEET1])
    return _pack(parts)


def _rewrite_parts(
    upload: ValidatedUpload, fn: Callable[[dict[str, bytes]], dict[str, bytes]]
) -> ValidatedUpload:
    return _pack(fn(_parts(upload)))


@pytest.fixture
def forbid_openpyxl(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    """被拒絕的檔案不得進 openpyxl：load_workbook 被呼叫就記錄下來讓測試失敗。"""
    calls: list[str] = []

    def record(*args: object, **kwargs: object) -> None:
        calls.append("load_workbook")
        raise AssertionError("惡意檔案不應交給 openpyxl 解析")

    monkeypatch.setattr(student_import_service, "load_workbook", record)
    return calls


_SST_REL = (
    b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
    b'relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>'
)
_SST_CT = (
    b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-'
    b'officedocument.spreadsheetml.sharedStrings+xml"/></Types>'
)


def _with_sst(parts: dict[str, bytes], sst_xml: bytes) -> dict[str, bytes]:
    parts["xl/sharedStrings.xml"] = sst_xml
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>", _SST_REL
    )
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(b"</Types>", _SST_CT)
    return parts


def _wide_row_whitespace(d: bytes) -> bytes:
    return d.replace(b"</row>", (b"<c\n/>" * 2000 + b'<c\tr="ZZ1"/>' * 100) + b"</row>", 1)


def _wide_row_ns_prefix(d: bytes) -> bytes:
    d = d.replace(
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
        b'<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
    )
    d = re.sub(
        rb"<(/?)(sheetPr|outlinePr|pageSetUpPr|dimension|sheetViews|sheetView|selection|"
        rb"sheetFormatPr|sheetData|row|c|v|is|t|pageMargins|worksheet)\b",
        rb"<\1x:\2",
        d,
    )
    return d.replace(b"</x:row>", b"<x:c/>" * 2000 + b"</x:row>", 1)


def _many_rows(d: bytes) -> bytes:
    rows = b"".join(b'<row r="%d"><c/></row>' % (n + 3) for n in range(MAX_IMPORT_ROWS + 20))
    return d.replace(b"</sheetData>", rows + b"</sheetData>", 1)


def _many_other_elems(d: bytes) -> bytes:
    return d.replace(b"<sheetData>", b"<sheetData>" + b"<a/>" * (MAX_XML_ELEMENTS + 1000), 1)


def _many_other_elems_root(d: bytes) -> bytes:
    return d.replace(b"</sheetData>", b"</sheetData>" + b"<a/>" * (MAX_XML_ELEMENTS + 1000), 1)


def _deep(d: bytes) -> bytes:
    return d.replace(b"</sheetData>", b"</sheetData>" + b"<a>" * 200 + b"</a>" * 200, 1)


def _long_attr(d: bytes) -> bytes:
    return d.replace(b"<sheetData>", b'<sheetData><row r="3" x="' + b"a" * 300_000 + b'"/>', 1)


def _long_text(d: bytes) -> bytes:
    cell = b'<row r="3"><c r="B3" t="inlineStr"><is><t>' + b"\xe7\x8e\x8b" * 100_000
    return d.replace(b"</sheetData>", cell + b"</t></is></c></row></sheetData>", 1)


_WIDE_ROW = b"<c/>" * 2000


def _no_decl(d: bytes) -> bytes:
    return re.sub(rb"^<\?xml[^>]*\?>", b"", d)


def _leading_ws(d: bytes) -> bytes:
    """無 XML 宣告 + 200 個前置空白：不能靠「開頭是否為 <」判斷是不是 XML。"""
    return b" " * 200 + _no_decl(d.replace(b"</row>", _WIDE_ROW + b"</row>", 1))


def _bom_ws(d: bytes) -> bytes:
    return b"\xef\xbb\xbf" + b"\n" * 100 + _no_decl(d.replace(b"</row>", _WIDE_ROW + b"</row>", 1))


def _utf16_sheet(d: bytes) -> bytes:
    x = d.replace(b"</row>", _WIDE_ROW + b"</row>", 1).decode("utf-8")
    x = re.sub(r"^<\?xml[^>]*\?>", '<?xml version="1.0" encoding="UTF-16"?>', x)
    return x.encode("utf-16")  # 含 BOM


def _utf16_nobom(d: bytes) -> bytes:
    x = d.replace(b"</row>", _WIDE_ROW + b"</row>", 1).decode("utf-8")
    x = re.sub(r"^<\?xml[^>]*\?>", '<?xml version="1.0" encoding="UTF-16"?>', x)
    return x.encode("utf-16-le")


def _utf32_sheet(d: bytes) -> bytes:
    x = re.sub(r"^<\?xml[^>]*\?>", '<?xml version="1.0" encoding="UTF-32"?>', d.decode("utf-8"))
    return x.encode("utf-32")


def _styles_bomb(parts: dict[str, bytes]) -> dict[str, bytes]:
    parts["xl/styles.xml"] = re.sub(
        rb'<cellXfs count="\d+">',
        b'<cellXfs count="1">'
        + b'<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        * (MAX_XML_ELEMENTS + 1000),
        parts["xl/styles.xml"],
        count=1,
    )
    return parts


def _sst_many(parts: dict[str, bytes]) -> dict[str, bytes]:
    sst = (
        b'<?xml version="1.0"?>'
        b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        + b"<si><t>a</t></si>" * (MAX_XML_ELEMENTS // 2 + 1000)
        + b"</sst>"
    )
    return _with_sst(parts, sst)


def _cap_styles(parts: dict[str, bytes]) -> dict[str, bytes]:
    """styles.xml 逼近舊的單 part 上限（25 萬）：合法檔只有數百個元素，必須被共用總額擋下。"""
    parts["xl/styles.xml"] = re.sub(
        rb'<cellXfs count="\d+">',
        b'<cellXfs count="1">'
        + b'<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>' * 249_000,
        parts["xl/styles.xml"],
        count=1,
    )
    return parts


def _cap_sst(parts: dict[str, bytes]) -> dict[str, bytes]:
    sst = (
        b'<?xml version="1.0"?>'
        b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        + b"<si><t>a</t></si>" * 124_000
        + b"</sst>"
    )
    return _with_sst(parts, sst)


def _combined_max(parts: dict[str, bytes]) -> dict[str, bytes]:
    """sheet / styles / sharedStrings 各自低於舊的單 part 上限，合計卻遠超合法檔：共用總額要擋。"""
    parts[_SHEET1] = parts[_SHEET1].replace(b"</sheetData>", b"</sheetData>" + b"<a/>" * 90_000, 1)
    parts = _cap_styles(parts)
    return _cap_sst(parts)


def _path_bypass(parts: dict[str, bytes]) -> dict[str, bytes]:
    """工作表改放到 worksheets/ 以外、非 .xml 副檔名，rels 與 Content_Types 同步改。"""
    sheet = parts.pop(_SHEET1).replace(b"</row>", b"<c/>" * 2000 + b"</row>", 1)
    parts["xl/data/s.bin"] = sheet
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"worksheets/sheet1.xml", b"data/s.bin"
    )
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"/xl/worksheets/sheet1.xml", b"/xl/data/s.bin"
    )
    return parts


def _many_entries(parts: dict[str, bytes]) -> dict[str, bytes]:
    for i in range(MAX_ZIP_ENTRIES + 1):
        parts[f"x/{i}"] = b""
    return parts


def _big_other_part(parts: dict[str, bytes]) -> dict[str, bytes]:
    parts["docProps/custom.xml"] = (
        b'<?xml version="1.0"?><Properties>'
        + b"<p/>" * (MAX_XML_ELEMENTS + 1000)
        + b"</Properties>"
    )
    return parts


def _lying_uncompressed_size(upload: ValidatedUpload) -> ValidatedUpload:
    """把 sheet1.xml 宣告的解壓大小改小（CRC / 大小不符）。"""
    c = bytearray(upload.content)
    name = _SHEET1.encode()
    for sig, size_off, nlen_off, fname_off in (
        (b"PK\x01\x02", 24, 28, 46),
        (b"PK\x03\x04", 22, 26, 30),
    ):
        i = 0
        while (i := c.find(sig, i)) >= 0:
            nlen = int.from_bytes(c[i + nlen_off : i + nlen_off + 2], "little")
            if bytes(c[i + fname_off : i + fname_off + nlen]) == name:
                c[i + size_off : i + size_off + 4] = (1000).to_bytes(4, "little")
            i += 4
    return ValidatedUpload(content=bytes(c), mime_type=upload.mime_type, ext="xlsx", size=len(c))


def _encrypted_flag(upload: ValidatedUpload) -> ValidatedUpload:
    """把 sheet1.xml 的 general purpose flag 加上 bit 0（加密）。"""
    c = bytearray(upload.content)
    name = _SHEET1.encode()
    for sig, flag_off, nlen_off, fname_off in (
        (b"PK\x01\x02", 8, 28, 46),
        (b"PK\x03\x04", 6, 26, 30),
    ):
        i = 0
        while (i := c.find(sig, i)) >= 0:
            nlen = int.from_bytes(c[i + nlen_off : i + nlen_off + 2], "little")
            if bytes(c[i + fname_off : i + fname_off + nlen]) == name:
                c[i + flag_off] |= 0x01
            i += 4
    return ValidatedUpload(content=bytes(c), mime_type=upload.mime_type, ext="xlsx", size=len(c))


def _duplicate_entry(upload: ValidatedUpload) -> ValidatedUpload:
    parts = _parts(upload)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in parts.items():
            dst.writestr(name, data)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)  # zipfile 對同名 entry 的提醒是刻意的
            dst.writestr(_SHEET1, parts[_SHEET1])  # 同名第二份
    content = out.getvalue()
    return ValidatedUpload(
        content=content, mime_type=upload.mime_type, ext="xlsx", size=len(content)
    )


def _bzip2(upload: ValidatedUpload) -> ValidatedUpload:
    return _pack(_parts(upload), compress=zipfile.ZIP_BZIP2)


def _truncated(upload: ValidatedUpload) -> ValidatedUpload:
    half = upload.content[: len(upload.content) // 2]
    return ValidatedUpload(content=half, mime_type=upload.mime_type, ext="xlsx", size=len(half))


_SHEET_ATTACKS: dict[str, Callable[[bytes], bytes]] = {
    "wide_ws": _wide_row_whitespace,
    "ns_prefix": _wide_row_ns_prefix,
    "many_other_elems": _many_other_elems,
    "many_other_elems_root": _many_other_elems_root,
    "deep": _deep,
    "long_attr": _long_attr,
    "long_text": _long_text,
    "leading_ws": _leading_ws,
    "bom_ws": _bom_ws,
    "utf16_sheet": _utf16_sheet,
    "utf16_nobom": _utf16_nobom,
    "utf32": _utf32_sheet,
}
_PARTS_ATTACKS: dict[str, Callable[[dict[str, bytes]], dict[str, bytes]]] = {
    "styles_bomb": _styles_bomb,
    "sst_many": _sst_many,
    "path_bypass": _path_bypass,
    "many_entries": _many_entries,
    "big_other_part": _big_other_part,
    "cap_styles": _cap_styles,
    "cap_sst": _cap_sst,
    "combined_max": _combined_max,
}
_ZIP_ATTACKS: dict[str, Callable[[ValidatedUpload], ValidatedUpload]] = {
    "lying_size": _lying_uncompressed_size,
    "encrypted": _encrypted_flag,
    "duplicate_entry": _duplicate_entry,
    "bzip2": _bzip2,
    "truncated": _truncated,
}


def _attack(name: str) -> ValidatedUpload:
    base = _xlsx([_row()], header=IMPORT_COLUMNS)
    if name in _SHEET_ATTACKS:
        return _rewrite_sheet(base, _SHEET_ATTACKS[name])
    if name in _PARTS_ATTACKS:
        return _rewrite_parts(base, _PARTS_ATTACKS[name])
    return _ZIP_ATTACKS[name](base)


@pytest.mark.parametrize("name", [*_SHEET_ATTACKS, *_PARTS_ATTACKS, *_ZIP_ATTACKS])
def test_import_preview_rejects_before_openpyxl(
    db_session: Session, forbid_openpyxl: list[str], name: str
) -> None:
    """每種構造檔都在 zip / XML 掃描階段以 422 拒絕、不進 openpyxl，且掃描本身不建樹
    （峰值 < 50 MB）。"""
    upload = _attack(name)
    assert upload.size < 2 * 1024 * 1024

    tracemalloc.start()
    try:
        with pytest.raises(AppError) as exc:
            _preview(db_session, upload)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert exc.value.status == 422
    assert exc.value.code == "import_invalid_file"
    assert forbid_openpyxl == []
    assert peak < 50 * 1024 * 1024, f"{name}: 掃描階段峰值 {peak / 1e6:.0f} MB"


def test_import_preview_openpyxl_parser_is_builtin_expat() -> None:
    """掃描器與 openpyxl 用同一個內建 expat 判斷「是不是 XML」；日後若安裝 lxml，接受範圍不同，
    需重新評估。"""
    assert openpyxl.xml.LXML is False


def test_import_preview_worst_case_legit_file_within_budget(
    db_session: Session, lookups: dict[str, object]
) -> None:
    """逼近業務上限的合法檔（500 列、13 欄、全部字串走 sharedStrings、每格都有樣式）仍被接受，
    且整個 preview 的 tracemalloc 峰值 < 100 MB。"""
    rows = [
        _row(
            student_no=f"S{n:06d}",
            name=f"學生{n}",
            school_class=f"二年{n % 9 + 1}班",
            note=f"備註 {n} " + "x" * 200,
        )
        for n in range(MAX_IMPORT_ROWS)
    ]
    upload = _xlsx(rows, header=IMPORT_COLUMNS)
    assert len(_parts(upload)) <= 12

    tracemalloc.start()
    try:
        result = _preview(db_session, upload)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert (result.total, result.valid) == (MAX_IMPORT_ROWS, MAX_IMPORT_ROWS)
    assert peak < 100 * 1024 * 1024, f"preview 峰值 {peak / 1e6:.0f} MB"


def test_import_preview_rejects_too_many_rows_before_openpyxl(
    db_session: Session, forbid_openpyxl: list[str]
) -> None:
    upload = _rewrite_sheet(_xlsx([_row()], header=IMPORT_COLUMNS), _many_rows)

    with pytest.raises(AppError) as exc:
        _preview(db_session, upload)

    assert (exc.value.status, exc.value.code) == (422, "import_too_many_rows")
    assert forbid_openpyxl == []


def test_import_preview_openpyxl_errors_become_422(db_session: Session) -> None:
    """掃描放行但 openpyxl 解析才出錯（sharedStrings 索引越界、壞的日期 / 數值）→ 422，不是 500。"""
    base = _xlsx([_row()], header=IMPORT_COLUMNS)
    big_index = _rewrite_sheet(
        base,
        lambda d: d.replace(
            b"</sheetData>",
            b'<row r="3"><c r="B3" t="s"><v>1000000000</v></c></row></sheetData>',
            1,
        ),
    )
    bad_date = _rewrite_sheet(
        base,
        lambda d: d.replace(
            b"</sheetData>",
            b'<row r="3"><c r="A3" t="d"><v>99999-99-99T99:99</v></c></row></sheetData>',
            1,
        ),
    )

    for upload in (big_index, bad_date):
        with pytest.raises(AppError) as exc:
            _preview(db_session, upload)
        assert (exc.value.status, exc.value.code) == _INVALID


def test_import_preview_uses_first_visible_sheet(
    db_session: Session, lookups: dict[str, object]
) -> None:
    """第一張工作表 hidden / veryHidden 時取第一張可見的；沒有任何可見工作表 → 422。"""
    wb = Workbook()
    hidden = wb.active
    assert hidden is not None
    hidden.title = "隱藏"
    hidden.sheet_state = "veryHidden"
    hidden.append(["垃圾"])
    visible = wb.create_sheet("學生資料")
    visible.append(list(IMPORT_COLUMNS))
    visible.append(_row(student_no="S115301"))
    buf = io.BytesIO()
    wb.save(buf)
    upload = ValidatedUpload(content=buf.getvalue(), mime_type="x", ext="xlsx", size=buf.tell())

    result = _preview(db_session, upload)

    assert result.valid == 1
    assert result.rows[0].data is not None
    assert result.rows[0].data.student_no == "S115301"

    # openpyxl 不允許存出全部隱藏的活頁簿：直接改 workbook.xml 把可見的那張也標 hidden
    all_hidden = _rewrite_parts(
        upload,
        lambda p: {
            **p,
            "xl/workbook.xml": p["xl/workbook.xml"].replace(
                "學生資料".encode(), "學生資料".encode() + b'" state="hidden', 1
            ),
        },
    )
    assert b'state="hidden' in _parts(all_hidden)["xl/workbook.xml"]
    with pytest.raises(AppError) as exc:
        _preview(db_session, all_hidden)
    assert (exc.value.status, exc.value.code) == _INVALID


def test_import_preview_accepts_normal_shared_strings_and_styles(
    db_session: Session, lookups: dict[str, object]
) -> None:
    """合理大小的 sharedStrings / styles / docProps 不受上限影響；OOXML _x0000_ 跳脫維持字面、
    不會 500。"""
    base = _xlsx([_row(), _row(student_no="S115302", name="陳小華")], header=IMPORT_COLUMNS)
    sst = (
        b'<?xml version="1.0"?>'
        b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        + b"<si><t>a</t></si>" * 2000
        + b"</sst>"
    )
    with_sst = _rewrite_parts(base, lambda p: _with_sst(p, sst))

    result = _preview(db_session, with_sst)
    assert result.valid == 2

    escaped = _rewrite_sheet(
        base,
        lambda d: d.replace(
            b"</sheetData>",
            b'<row r="4"><c r="A4" t="inlineStr"><is><t>S115303</t></is></c>'
            + b'<c r="B4" t="inlineStr"><is><t>'
            + "王_x0000_明".encode()
            + b"</t></is></c>"
            + b'<c r="E4"><v>2</v></c></row></sheetData>',
            1,
        ),
    )
    escaped_result = _preview(db_session, escaped)
    # openpyxl 不還原 _xHHHH_ 跳脫：維持字面字串（不會變成 NUL 讓 DB 寫入失敗），也不會 500
    assert escaped_result.total == 3
    assert escaped_result.rows[2].data is not None
    assert escaped_result.rows[2].data.name == "王_x0000_明"
    assert "\x00" not in escaped_result.rows[2].data.name


# --- BACKEND-156：execute -------------------------------------------------------------------------


def _execute(
    db: Session, upload: ValidatedUpload, clock: FakeClock, actor: CurrentStaff = _SENSITIVE
) -> ImportResult:
    return execute(db, upload, academic_year=115, actor=actor, meta=_META, clock=clock)


def _by_student_no(db: Session, *student_nos: str) -> list[Student]:
    return list(
        db.execute(
            select(Student).where(Student.student_no.in_(student_nos)).order_by(Student.student_no)
        ).scalars()
    )


def _import_audits(db: Session) -> list[AuditLog]:
    return list(db.execute(select(AuditLog).where(AuditLog.action == "student.import")).scalars())


def _sensitive_audits(db: Session, student_id: object) -> list[AuditLog]:
    return list(
        db.execute(
            select(AuditLog).where(
                AuditLog.action == "student.sensitive_update",
                AuditLog.entity_id == str(student_id),
            )
        ).scalars()
    )


def test_import_execute_success(
    db_session: Session,
    db_engine: Engine,
    lookups: dict[str, object],
    fake_clock: FakeClock,
) -> None:
    upload = _xlsx(
        [
            _row(),
            _row(student_no="S115102", name="陳小華", gender="女", status="暫停"),
            _row(student_no="S115103", name="黃小美", school=None, klass=None, status="退班"),
        ],
        header=IMPORT_COLUMNS,
    )
    actor = _SENSITIVE

    result = _execute(db_session, upload, fake_clock, actor=actor)

    assert isinstance(result, ImportResult)
    assert result.created == 3
    assert len(result.student_ids) == 3
    students = _by_student_no(db_session, "S115101", "S115102", "S115103")
    assert [s.student_no for s in students] == ["S115101", "S115102", "S115103"]
    assert sorted(result.student_ids) == sorted(s.id for s in students)
    first, second, third = students
    assert (first.name, first.gender, first.grade_level, first.status) == (
        "林小安",
        "male",
        2,
        "active",
    )
    assert first.class_id == lookups["class"].id  # type: ignore[attr-defined]
    assert first.school_id == lookups["school"].id  # type: ignore[attr-defined]
    assert (first.school_class, first.enrolled_on) == ("二年三班", date(2026, 9, 1))
    assert (first.id_number_enc, first.id_number_hmac, first.health_note_enc) == (None, None, None)
    assert (second.gender, second.status) == ("female", "suspended")
    # 退班未給退班日 → 以 clock 的今天補（同 create_student）
    assert (third.status, third.withdrawn_on, third.class_id) == (
        "withdrawn",
        fake_clock.today(),
        None,
    )
    audits = _import_audits(db_session)
    assert len(audits) == 1
    assert audits[0].after == {"created": 3, "academic_year": 115}
    assert audits[0].entity_type == "student_import"
    assert (audits[0].actor_type, audits[0].actor_id, audits[0].ip) == (
        "staff",
        actor.id,
        "127.0.0.1",
    )
    # 不 commit：另一條連線看不到
    with Session(bind=db_engine) as other:
        assert (
            other.execute(
                select(func.count())
                .select_from(Student)
                .where(Student.student_no.in_(["S115101", "S115102", "S115103"]))
            ).scalar_one()
            == 0
        )


def test_import_execute_has_errors(
    db_session: Session, lookups: dict[str, object], fake_clock: FakeClock
) -> None:
    before = db_session.execute(select(func.count()).select_from(Student)).scalar_one()
    upload = _xlsx(
        [
            _row(),
            _row(student_no="S115102", name="陳小華", school="不存在國小"),
            _row(student_no="S115103", name="黃小美"),
        ],
        header=IMPORT_COLUMNS,
    )

    with pytest.raises(AppError) as exc:
        _execute(db_session, upload, fake_clock)

    assert (exc.value.status, exc.value.code) == (409, "import_has_errors")
    details = exc.value.details
    assert (details["total"], details["valid"], details["invalid"]) == (3, 2, 1)
    assert len(details["rows"]) == 1  # 只含有錯誤的列
    bad = details["rows"][0]
    assert bad["row_number"] == 3
    assert bad["errors"] == ["找不到國小：不存在國小"]
    assert bad["display"]["學號*"] == "S115102"
    assert bad["display"]["就讀國小"] == "不存在國小"
    assert "data" not in bad  # 正規化資料（可能含敏感明文）不回給前端
    assert db_session.execute(select(func.count()).select_from(Student)).scalar_one() == before
    assert _by_student_no(db_session, "S115101", "S115103") == []
    assert _import_audits(db_session) == []
    # 檔案層錯誤原樣往上拋、不寫入
    with pytest.raises(AppError) as empty:
        _execute(db_session, _xlsx([], header=IMPORT_COLUMNS), fake_clock)
    assert (empty.value.status, empty.value.code) == (422, "import_empty")
    assert db_session.execute(select(func.count()).select_from(Student)).scalar_one() == before


def test_import_execute_encrypts(
    db_session: Session, lookups: dict[str, object], fake_clock: FakeClock
) -> None:
    upload = _xlsx(
        [
            _row(id_number=f" {_ID_A.lower()} ", health_note="氣喘"),
            _row(student_no="S115102", name="陳小華"),
        ],
        header=IMPORT_COLUMNS,
    )

    result = _execute(db_session, upload, fake_clock)

    assert result.created == 2
    first, second = _by_student_no(db_session, "S115101", "S115102")
    assert first.id_number_enc is not None
    assert _ID_A.encode() not in first.id_number_enc
    assert decrypt_bytes(first.id_number_enc) == _ID_A
    assert first.id_number_hmac == id_number_hmac(_ID_A)
    assert first.health_note_enc is not None
    assert "氣喘".encode() not in first.health_note_enc
    assert decrypt_bytes(first.health_note_enc) == "氣喘"
    assert (second.id_number_enc, second.id_number_hmac, second.health_note_enc) == (
        None,
        None,
        None,
    )
    # 與 create_student 相同：每位有敏感欄位的學生各一筆 sensitive_update 稽核（只記欄位名）
    assert [a.after for a in _sensitive_audits(db_session, first.id)] == [
        {"set": ["health_note", "id_number"]}
    ]
    assert _sensitive_audits(db_session, second.id) == []
    # 無 students:sensitive 的匯入者：preview 就把帶敏感欄位的列標成錯誤 → 409、不寫入
    with pytest.raises(AppError) as exc:
        _execute(
            db_session,
            _xlsx([_row(student_no="S115103", id_number=_ID_B)], header=IMPORT_COLUMNS),
            fake_clock,
            actor=_WRITER,
        )
    assert exc.value.code == "import_has_errors"
    assert "沒有權限匯入敏感欄位" in exc.value.details["rows"][0]["errors"]
    assert _by_student_no(db_session, "S115103") == []


def test_import_execute_conflict_rolls_back_all(
    db_session: Session,
    lookups: dict[str, object],
    fake_clock: FakeClock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """preview 全部合法，寫入前另一筆同學號被人插入（競態）→ 第 2 列撞 unique → 整批回滾、
    409 import_conflict，第 1 / 3 列也不存在，session 仍可用。"""
    real_preview = student_import_service.preview

    def preview_then_race(*args: object, **kwargs: object) -> ImportPreview:
        result = real_preview(*args, **kwargs)  # type: ignore[arg-type]
        make_student(db_session, name="搶先者", student_no="S115102")
        return result

    monkeypatch.setattr(student_import_service, "preview", preview_then_race)
    upload = _xlsx(
        [
            _row(id_number=_ID_A),
            _row(student_no="S115102", name="陳小華"),
            _row(student_no="S115103", name="黃小美"),
        ],
        header=IMPORT_COLUMNS,
    )

    with pytest.raises(AppError) as exc:
        _execute(db_session, upload, fake_clock)

    assert (exc.value.status, exc.value.code) == (409, "import_conflict")
    assert exc.value.details == {
        "row_number": 3,
        "code": "student_no_taken",
        "message": "學號已被使用",
    }
    students = _by_student_no(db_session, "S115101", "S115102", "S115103")
    assert [(s.student_no, s.name) for s in students] == [("S115102", "搶先者")]
    assert _import_audits(db_session) == []
    assert (
        db_session.execute(
            select(func.count())
            .select_from(AuditLog)
            .where(AuditLog.action == "student.sensitive_update")
        ).scalar_one()
        == 0
    )
    # savepoint 已回滾：同一 session 修掉衝突列後可以成功匯入
    monkeypatch.setattr(student_import_service, "preview", real_preview)
    fixed = _xlsx(
        [_row(id_number=_ID_A), _row(student_no="S115103", name="黃小美")], header=IMPORT_COLUMNS
    )
    assert _execute(db_session, fixed, fake_clock).created == 2
