"""BACKEND-159：GET /api/admin/students。
BACKEND-161：GET /api/admin/students/{student_id}（students:read；sensitive 另需
students:sensitive）。
BACKEND-160：POST /api/admin/students（students:write；敏感欄位另需 students:sensitive）。
BACKEND-163：POST /api/admin/students/{student_id}/archive（封存後家長端看不到）。
BACKEND-531：POST /api/admin/students/{student_id}/purge（students:purge，永久刪除 = 匿名化）。
BACKEND-164：POST /api/admin/students/{student_id}/photo（multipart file，students:write）。
BACKEND-162：PATCH /api/admin/students/{student_id}（students:write；回應以 get_student 重組、簽
photo_url；改狀態同交易收尾並 commit）。
BACKEND-533：GET /api/admin/students/import-template（students:write；xlsx 範本、固定檔名、
no-store、路由在 /{student_id} 之前）。
"""

from __future__ import annotations

from collections.abc import Callable
from datetime import UTC, date, datetime
from io import BytesIO
from urllib.parse import unquote
from uuid import uuid4

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.crypto import decrypt_bytes, encrypt_bytes
from app.models.account import StaffUser
from app.models.audit import AuditLog
from app.models.parents import ParentAccount
from app.models.pickup import PickupRequest
from app.models.students import Student
from app.services.student_import_service import IMPORT_COLUMNS
from app.services.students.id_number import id_number_hmac
from tests.support.factories import (
    make_class,
    make_guardian,
    make_pickup_request,
    make_school,
    make_student,
)
from tests.support.fake_clock import FakeClock
from tests.support.fake_storage import FakeStorage
from tests.support.route_audit import admin_routes_without_permission

_URL = "/api/admin/students"
StaffClientFactory = Callable[..., tuple[TestClient, StaffUser]]
ParentClientFactory = Callable[..., tuple[TestClient, ParentAccount]]
AssertError = Callable[..., None]
_ID_NUMBER = "A123456789"


def test_admin_students_list_success(staff_client: StaffClientFactory, db_session: Session) -> None:
    klass = make_class(db_session, name="三年甲班")
    school = make_school(db_session, name="新生國小")
    make_student(db_session, name="王小明", student_no="S115001", class_=klass, school=school)
    make_student(db_session, name="陳小華", student_no="S115002")
    client, _ = staff_client(permissions=["students:read"])

    resp = client.get(_URL, params={"q": "王"})

    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 1
    item = body["items"][0]
    assert item["name"] == "王小明"
    assert item["student_no"] == "S115001"
    assert item["class"]["name"] == "三年甲班"
    assert "class_" not in item
    assert item["school"]["name"] == "新生國小"
    assert set(item) == {
        "id",
        "student_no",
        "name",
        "gender",
        "grade_level",
        "school",
        "school_class",
        "class",
        "status",
        "archived_at",
    }


def test_admin_students_list_filters_and_paging(
    staff_client: StaffClientFactory, db_session: Session
) -> None:
    klass = make_class(db_session)
    for n in range(25):
        make_student(db_session, class_=klass, student_no=f"P{n:03d}", grade_level=3)
    make_student(db_session, class_=klass, student_no="ARCH", archived=True)
    client, _ = staff_client(permissions=["students:read"])

    page3 = client.get(_URL, params={"class_id": str(klass.id), "page": 3, "page_size": 10}).json()
    with_archived = client.get(
        _URL, params={"class_id": str(klass.id), "include_archived": "true", "page_size": 100}
    ).json()

    assert page3["total"] == 25
    assert [i["student_no"] for i in page3["items"]] == [f"P{n:03d}" for n in range(20, 25)]
    assert with_archived["total"] == 26
    assert (
        client.get(_URL, params={"class_id": str(klass.id), "grade_level": 4}).json()["total"] == 0
    )


def test_admin_students_list_422(
    staff_client: StaffClientFactory, assert_error: AssertError
) -> None:
    client, _ = staff_client(permissions=["students:read"])

    for params in (
        {"grade_level": 9},
        {"status": "graduated"},
        {"class_id": "not-a-uuid"},
        {"q": "x" * 51},
        {"include_archived": "maybe"},
        {"page": 0},
        {"unknown": "1"},
    ):
        assert_error(client.get(_URL, params=params), 422, "validation_error")


def test_admin_students_list_401(api_client: TestClient, assert_error: AssertError) -> None:
    assert_error(api_client.get(_URL), 401, "unauthenticated")


def test_admin_students_list_403(
    staff_client: StaffClientFactory, assert_error: AssertError
) -> None:
    client, _ = staff_client(permissions=["classes:read"])

    resp = client.get(_URL)

    assert_error(resp, 403, "permission_denied")
    assert resp.json()["error"]["details"] == {"required": ["students:read"]}


def test_admin_students_list_no_sensitive(
    staff_client: StaffClientFactory, db_session: Session
) -> None:
    student = make_student(db_session, name="王小明")
    student.id_number_enc = b"\x01encrypted-id"
    student.id_number_hmac = "a" * 64
    student.health_note_enc = b"\x01encrypted-health"
    db_session.flush()
    client, _ = staff_client(permissions=["students:read", "students:sensitive"])

    resp = client.get(_URL, params={"q": "王小明"})

    assert resp.status_code == 200
    for leaked in (_ID_NUMBER, "id_number", "health_note", "encrypted", "hmac", "sensitive"):
        assert leaked not in resp.text


def test_admin_students_list_guard_registered(app: FastAPI) -> None:
    assert admin_routes_without_permission(app) == []
    assert "/api/admin/students" in app.openapi()["paths"]


# --- BACKEND-161：GET /api/admin/students/{student_id} -----------------------------


def _with_sensitive(db: Session) -> Student:
    student = make_student(db, name="王小明", student_no="S115001")
    student.id_number_enc = encrypt_bytes(_ID_NUMBER)
    student.id_number_hmac = id_number_hmac(_ID_NUMBER)
    student.health_note_enc = encrypt_bytes("對花生過敏")
    student.photo_path = f"{student.id}/{'ab' * 16}.jpg"
    make_guardian(db, student, name="王媽媽", is_primary=True)
    db.commit()
    return student


def test_admin_students_get_sensitive(
    staff_client: StaffClientFactory, db_session: Session
) -> None:
    student = _with_sensitive(db_session)
    client, _ = staff_client(permissions=["students:read", "students:sensitive"])

    resp = client.get(f"{_URL}/{student.id}")

    assert resp.status_code == 200
    body = resp.json()
    assert body["id"] == str(student.id)
    assert body["name"] == "王小明"
    assert body["sensitive"] == {"id_number": _ID_NUMBER, "health_note": "對花生過敏"}
    assert (body["has_id_number"], body["has_health_note"]) == (True, True)
    assert body["photo_url"] == (
        f"https://storage.test/student-photos/{student.id}/{'ab' * 16}.jpg?exp=300"
    )
    assert [g["name"] for g in body["guardians"]] == ["王媽媽"]
    assert body["class"] is None
    assert "class_" not in body
    assert set(body) == {
        "id",
        "student_no",
        "name",
        "gender",
        "grade_level",
        "school",
        "school_class",
        "class",
        "status",
        "archived_at",
        "birthday",
        "enrolled_on",
        "withdrawn_on",
        "note",
        "photo_url",
        "has_id_number",
        "has_health_note",
        "sensitive",
        "guardians",
    }


def test_admin_students_get_hidden(staff_client: StaffClientFactory, db_session: Session) -> None:
    student = _with_sensitive(db_session)
    archived = make_student(db_session, archived=True)
    db_session.commit()
    client, _ = staff_client(permissions=["students:read"])

    resp = client.get(f"{_URL}/{student.id}")

    assert resp.status_code == 200
    body = resp.json()
    assert body["sensitive"] is None
    assert (body["has_id_number"], body["has_health_note"]) == (True, True)
    for leaked in (_ID_NUMBER, "對花生過敏", "_enc", "hmac"):
        assert leaked not in resp.text
    # 封存學生也可查
    archived_resp = client.get(f"{_URL}/{archived.id}")
    assert archived_resp.status_code == 200
    assert archived_resp.json()["archived_at"] is not None


def test_admin_students_get_422_401(
    staff_client: StaffClientFactory, api_client: TestClient, assert_error: AssertError
) -> None:
    client, _ = staff_client(permissions=["students:read"])

    assert_error(client.get(f"{_URL}/abc"), 422, "validation_error")
    assert_error(api_client.get(f"{_URL}/{uuid4()}"), 401, "unauthenticated")


def test_admin_students_get_403(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    student = make_student(db_session)
    client, _ = staff_client(permissions=["classes:read"])

    resp = client.get(f"{_URL}/{student.id}")

    assert_error(resp, 403, "permission_denied")
    assert resp.json()["error"]["details"] == {"required": ["students:read"]}


def test_admin_students_get_404(
    staff_client: StaffClientFactory, assert_error: AssertError
) -> None:
    client, _ = staff_client(permissions=["students:read", "students:sensitive"])

    assert_error(client.get(f"{_URL}/{uuid4()}"), 404, "student_not_found")


def test_admin_students_get_guard_registered(app: FastAPI) -> None:
    assert admin_routes_without_permission(app) == []
    assert "get" in app.openapi()["paths"]["/api/admin/students/{student_id}"]


# --- BACKEND-160：POST /api/admin/students -----------------------------------------

_CREATE = {"student_no": "S115020", "name": "林小安", "grade_level": 2}


def test_admin_students_create_success(
    staff_client: StaffClientFactory, db_session: Session
) -> None:
    client, _ = staff_client(permissions=["students:write", "students:read"])

    resp = client.post(_URL, json=_CREATE)

    assert resp.status_code == 201
    body = resp.json()
    assert body["name"] == "林小安"
    assert body["student_no"] == "S115020"
    assert body["sensitive"] is None  # 沒有 students:sensitive
    assert body["has_id_number"] is False
    assert body["status"] == "active"
    assert body["guardians"] == []
    stored = db_session.execute(select(Student).where(Student.student_no == "S115020")).scalar_one()
    assert (stored.name, stored.grade_level, stored.id_number_enc) == ("林小安", 2, None)
    # 有 sensitive 權限時可寫身分證，回應含明文、DB 只有密文
    sensitive_client, _ = staff_client(permissions=["students:write", "students:sensitive"])
    created = sensitive_client.post(
        _URL, json={**_CREATE, "student_no": "S115021", "id_number": " a123456789 "}
    )
    assert created.status_code == 201
    assert created.json()["sensitive"]["id_number"] == _ID_NUMBER
    with_id = db_session.execute(
        select(Student).where(Student.student_no == "S115021")
    ).scalar_one()
    assert with_id.id_number_enc is not None
    assert decrypt_bytes(with_id.id_number_enc) == _ID_NUMBER
    assert with_id.id_number_hmac == id_number_hmac(_ID_NUMBER)


def test_admin_students_create_422(
    staff_client: StaffClientFactory, assert_error: AssertError
) -> None:
    client, _ = staff_client(permissions=["students:write", "students:sensitive"])

    assert_error(client.post(_URL, json={**_CREATE, "grade_level": 0}), 422, "validation_error")
    assert_error(client.post(_URL, json={**_CREATE, "id_number": "A123"}), 422, "invalid_id_number")
    assert_error(client.post(_URL, json={**_CREATE, "foo": 1}), 422, "validation_error")
    assert_error(client.post(_URL, json={}), 422, "validation_error")


def test_admin_students_create_401(api_client: TestClient, assert_error: AssertError) -> None:
    assert_error(api_client.post(_URL, json=_CREATE), 401, "unauthenticated")


def test_admin_students_create_403(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    reader, _ = staff_client(permissions=["students:read"])
    writer, _ = staff_client(permissions=["students:write"])

    denied = reader.post(_URL, json=_CREATE)
    sensitive = writer.post(_URL, json={**_CREATE, "id_number": _ID_NUMBER})

    assert_error(denied, 403, "permission_denied")
    assert denied.json()["error"]["details"] == {"required": ["students:write"]}
    assert_error(sensitive, 403, "sensitive_permission_required")
    assert (
        db_session.execute(
            select(Student).where(Student.student_no == "S115020")
        ).scalar_one_or_none()
        is None
    )


def test_admin_students_create_409(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    make_student(db_session, student_no="S115020")
    db_session.commit()
    client, _ = staff_client(permissions=["students:write"])

    assert_error(client.post(_URL, json=_CREATE), 409, "student_no_taken")


def test_admin_students_create_guard_registered(app: FastAPI) -> None:
    assert admin_routes_without_permission(app) == []
    assert "post" in app.openapi()["paths"]["/api/admin/students"]


# --- BACKEND-163：POST /api/admin/students/{id}/archive ----------------------------


def _archive_url(student_id: object) -> str:
    return f"{_URL}/{student_id}/archive"


def test_admin_students_archive_success(
    staff_client: StaffClientFactory, db_session: Session
) -> None:
    student = make_student(db_session, name="王小明")
    client, _ = staff_client(permissions=["students:write"])

    resp = client.post(_archive_url(student.id))

    assert resp.status_code == 200
    body = resp.json()
    assert body["id"] == str(student.id)
    assert body["archived_at"] is not None
    assert body["name"] == "王小明"
    db_session.expire_all()
    assert student.archived_at is not None
    # 冪等
    again = client.post(_archive_url(student.id))
    assert again.status_code == 200
    assert again.json()["archived_at"] == body["archived_at"]


def test_admin_students_archive_parent_hidden(
    staff_client: StaffClientFactory, parent_client: ParentClientFactory, db_session: Session
) -> None:
    parent, parent_account = parent_client()
    student = make_student(db_session, name="王小明")
    make_guardian(db_session, student, parent=parent_account)
    db_session.commit()
    assert [c["name"] for c in parent.get("/api/parent/children").json()] == ["王小明"]
    client, _ = staff_client(permissions=["students:write"])

    assert client.post(_archive_url(student.id)).status_code == 200

    assert parent.get("/api/parent/children").json() == []


def test_admin_students_archive_422_401(
    staff_client: StaffClientFactory, api_client: TestClient, assert_error: AssertError
) -> None:
    client, _ = staff_client(permissions=["students:write"])

    assert_error(client.post(_archive_url("abc")), 422, "validation_error")
    assert_error(api_client.post(_archive_url(uuid4())), 401, "unauthenticated")


def test_admin_students_archive_403(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    student = make_student(db_session)
    client, _ = staff_client(permissions=["students:read"])

    resp = client.post(_archive_url(student.id))

    assert_error(resp, 403, "permission_denied")
    db_session.expire_all()
    assert student.archived_at is None


def test_admin_students_archive_404(
    staff_client: StaffClientFactory, assert_error: AssertError
) -> None:
    client, _ = staff_client(permissions=["students:write"])

    assert_error(client.post(_archive_url(uuid4())), 404, "student_not_found")


# --- BACKEND-531：POST /api/admin/students/{id}/purge ------------------------------


def _purge_url(student_id: object) -> str:
    return f"{_URL}/{student_id}/purge"


def _purgeable(db: Session) -> Student:
    student = make_student(db, name="王小明", student_no="S115001")
    student.status = "withdrawn"
    student.withdrawn_on = date(2026, 8, 31)
    student.archived_at = datetime(2026, 9, 1, tzinfo=UTC)
    student.id_number_enc = encrypt_bytes(_ID_NUMBER)
    student.id_number_hmac = id_number_hmac(_ID_NUMBER)
    db.commit()
    return student


def test_admin_students_purge_success(
    staff_client: StaffClientFactory, db_session: Session
) -> None:
    student = _purgeable(db_session)
    client, staff = staff_client(role_code="admin")

    resp = client.post(_purge_url(student.id), json={"confirm_student_no": "S115001"})

    assert resp.status_code == 200
    body = resp.json()
    assert body["student_id"] == str(student.id)
    assert body["anonymized_student_no"].startswith("DEL-")
    assert body["purged_at"] is not None
    detail = client.get(f"{_URL}/{student.id}").json()
    assert detail["name"] == "已刪除學生"
    assert detail["student_no"] == body["anonymized_student_no"]
    assert detail["has_id_number"] is False
    db_session.expire_all()
    assert (student.name, student.id_number_enc, student.id_number_hmac) == (
        "已刪除學生",
        None,
        None,
    )
    audits = (
        db_session.execute(
            select(AuditLog).where(
                AuditLog.action == "student.purge", AuditLog.entity_id == str(student.id)
            )
        )
        .scalars()
        .all()
    )
    assert len(audits) == 1
    assert audits[0].actor_id == staff.id


def test_admin_students_purge_422(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    student = _purgeable(db_session)
    client, _ = staff_client(role_code="admin")

    assert_error(client.post(_purge_url(student.id), json={}), 422, "validation_error")
    assert_error(
        client.post(_purge_url(student.id), json={"confirm_student_no": "S999999"}),
        422,
        "purge_confirmation_mismatch",
    )
    db_session.expire_all()
    assert student.name == "王小明"


def test_admin_students_purge_401(api_client: TestClient, assert_error: AssertError) -> None:
    assert_error(
        api_client.post(_purge_url(uuid4()), json={"confirm_student_no": "S115001"}),
        401,
        "unauthenticated",
    )


def test_admin_students_purge_403(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    student = _purgeable(db_session)
    director, _ = staff_client(role_code="director")

    resp = director.post(_purge_url(student.id), json={"confirm_student_no": "S115001"})

    assert_error(resp, 403, "permission_denied")
    assert resp.json()["error"]["details"] == {"required": ["students:purge"]}
    db_session.expire_all()
    assert student.name == "王小明"


def test_admin_students_purge_409(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    active = make_student(db_session, student_no="ACT001")
    student = _purgeable(db_session)
    client, _ = staff_client(role_code="admin")

    assert_error(
        client.post(_purge_url(active.id), json={"confirm_student_no": "ACT001"}),
        409,
        "student_not_purgeable",
    )
    first = client.post(_purge_url(student.id), json={"confirm_student_no": "S115001"})
    assert first.status_code == 200
    new_no = first.json()["anonymized_student_no"]
    assert_error(
        client.post(_purge_url(student.id), json={"confirm_student_no": new_no}),
        409,
        "student_already_purged",
    )


def test_admin_students_purge_guard_registered(app: FastAPI) -> None:
    assert admin_routes_without_permission(app) == []
    assert "post" in app.openapi()["paths"]["/api/admin/students/{student_id}/purge"]


# --- BACKEND-164：POST /api/admin/students/{id}/photo ------------------------------

_JPEG = b"\xff\xd8\xff\xe0" + b"0" * 100
_PDF = b"%PDF-1.7\n" + b"0" * 50


def _photo_url(student_id: object) -> str:
    return f"{_URL}/{student_id}/photo"


def test_admin_students_photo_success(
    staff_client: StaffClientFactory, db_session: Session, fake_storage: FakeStorage
) -> None:
    student = make_student(db_session)
    client, _ = staff_client(permissions=["students:write"])

    resp = client.post(_photo_url(student.id), files={"file": ("a.jpg", _JPEG, "image/jpeg")})

    assert resp.status_code == 200
    body = resp.json()
    assert set(body) == {"photo_url"}
    assert body["photo_url"].startswith("https://storage.test/student-photos/")
    db_session.expire_all()
    assert student.photo_path is not None
    assert body["photo_url"] == f"https://storage.test/student-photos/{student.photo_path}?exp=300"
    assert fake_storage.objects == {("student-photos", student.photo_path): _JPEG}


def test_admin_students_photo_422(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    student = make_student(db_session)
    client, _ = staff_client(permissions=["students:write"])

    assert_error(client.post(_photo_url(student.id)), 422, "validation_error")
    assert_error(
        client.post(_photo_url("abc"), files={"file": ("a.jpg", _JPEG, "image/jpeg")}),
        422,
        "validation_error",
    )


def test_admin_students_photo_401(api_client: TestClient, assert_error: AssertError) -> None:
    resp = api_client.post(_photo_url(uuid4()), files={"file": ("a.jpg", _JPEG, "image/jpeg")})

    assert_error(resp, 401, "unauthenticated")


def test_admin_students_photo_403(
    staff_client: StaffClientFactory,
    assert_error: AssertError,
    db_session: Session,
    fake_storage: FakeStorage,
) -> None:
    student = make_student(db_session)
    client, _ = staff_client(permissions=["students:read"])

    resp = client.post(_photo_url(student.id), files={"file": ("a.jpg", _JPEG, "image/jpeg")})

    assert_error(resp, 403, "permission_denied")
    assert resp.json()["error"]["details"] == {"required": ["students:write"]}
    assert fake_storage.objects == {}


def test_admin_students_photo_415(
    staff_client: StaffClientFactory,
    assert_error: AssertError,
    db_session: Session,
    fake_storage: FakeStorage,
) -> None:
    student = make_student(db_session)
    archived = make_student(db_session, archived=True)
    client, _ = staff_client(permissions=["students:write"])

    pdf = client.post(_photo_url(student.id), files={"file": ("a.pdf", _PDF, "application/pdf")})
    missing = client.post(_photo_url(uuid4()), files={"file": ("a.jpg", _JPEG, "image/jpeg")})
    gone = client.post(_photo_url(archived.id), files={"file": ("a.jpg", _JPEG, "image/jpeg")})

    assert_error(pdf, 415, "unsupported_file_type")
    assert_error(missing, 404, "student_not_found")
    assert_error(gone, 404, "student_not_found")
    assert fake_storage.objects == {}
    db_session.expire_all()
    assert student.photo_path is None


def test_admin_students_photo_guard_registered(app: FastAPI) -> None:
    assert admin_routes_without_permission(app) == []
    assert "post" in app.openapi()["paths"]["/api/admin/students/{student_id}/photo"]


# --- BACKEND-162：PATCH /api/admin/students/{id} ------------------------------------


def _student_url(student_id: object) -> str:
    return f"{_URL}/{student_id}"


def _stored(db: Session, student_id: object) -> Student:
    db.expire_all()
    return db.execute(select(Student).where(Student.id == student_id)).scalar_one()


def _audits(db: Session, action: str, entity_id: object) -> list[AuditLog]:
    return list(
        db.execute(
            select(AuditLog).where(AuditLog.action == action, AuditLog.entity_id == str(entity_id))
        ).scalars()
    )


def test_admin_students_update_success(
    staff_client: StaffClientFactory, db_session: Session, fake_clock: FakeClock
) -> None:
    student = make_student(db_session, name="王小明", grade_level=3)
    pending = make_pickup_request(db_session, student, service_date=fake_clock.today())
    db_session.commit()
    client, staff = staff_client(permissions=["students:write"])

    resp = client.patch(_student_url(student.id), json={"status": "withdrawn"})

    assert resp.status_code == 200
    body = resp.json()
    assert body["id"] == str(student.id)
    assert (body["status"], body["withdrawn_on"]) == ("withdrawn", fake_clock.today().isoformat())
    assert (body["name"], body["grade_level"]) == ("王小明", 3)  # 只動給定欄位
    assert body["sensitive"] is None
    assert body["photo_url"] is None
    # 已 commit：重讀 DB 看得到狀態、529 收尾（接送請求取消）與稽核
    stored = _stored(db_session, student.id)
    assert (stored.status, stored.withdrawn_on) == ("withdrawn", fake_clock.today())
    assert db_session.get(PickupRequest, pending.id).status == "cancelled"  # type: ignore[union-attr]
    close_outs = _audits(db_session, "student.close_out", student.id)
    assert len(close_outs) == 1
    assert close_outs[0].actor_id == staff.id
    # 部分更新：改回 active 清除退班日，再只改 note 其他欄位不動
    reactivated = client.patch(_student_url(student.id), json={"status": "active"})
    assert (reactivated.json()["status"], reactivated.json()["withdrawn_on"]) == ("active", None)
    noted = client.patch(_student_url(student.id), json={"note": "喜歡數學"})
    assert (noted.json()["note"], noted.json()["status"]) == ("喜歡數學", "active")
    assert _stored(db_session, student.id).note == "喜歡數學"
    assert len(_audits(db_session, "student.close_out", student.id)) == 1


def test_admin_students_update_photo_url(
    staff_client: StaffClientFactory, db_session: Session
) -> None:
    with_photo = make_student(db_session, name="王小明")
    with_photo.photo_path = f"{with_photo.id}/{'ab' * 16}.jpg"
    without_photo = make_student(db_session, name="陳小華")
    db_session.commit()
    client, _ = staff_client(permissions=["students:write", "students:read"])

    resp = client.patch(_student_url(with_photo.id), json={"note": "x"})
    plain = client.patch(_student_url(without_photo.id), json={"note": "y"})

    assert resp.status_code == 200
    assert resp.json()["photo_url"] == (
        f"https://storage.test/student-photos/{with_photo.id}/{'ab' * 16}.jpg?exp=300"
    )
    # 與 GET 詳情同一個組裝函式：兩者 photo_url 一致
    assert resp.json()["photo_url"] == client.get(_student_url(with_photo.id)).json()["photo_url"]
    assert plain.status_code == 200
    assert plain.json()["photo_url"] is None
    assert set(resp.json()) == set(client.get(_student_url(with_photo.id)).json())


def test_admin_students_update_sensitive(
    staff_client: StaffClientFactory, db_session: Session
) -> None:
    student = _with_sensitive(db_session)
    client, staff = staff_client(permissions=["students:write", "students:sensitive"])

    resp = client.patch(_student_url(student.id), json={"id_number": None, "health_note": "氣喘"})

    assert resp.status_code == 200
    body = resp.json()
    assert body["sensitive"] == {"id_number": None, "health_note": "氣喘"}
    assert (body["has_id_number"], body["has_health_note"]) == (False, True)
    stored = _stored(db_session, student.id)
    assert (stored.id_number_enc, stored.id_number_hmac) == (None, None)
    assert stored.health_note_enc is not None
    assert decrypt_bytes(stored.health_note_enc) == "氣喘"
    audits = _audits(db_session, "student.sensitive_update", student.id)
    assert [a.after for a in audits] == [{"cleared": ["id_number"], "set": ["health_note"]}]
    assert audits[0].actor_id == staff.id
    # 只有 students:write：回應 sensitive 為 None、本文不含明文（has_* 仍回）
    writer, _ = staff_client(permissions=["students:write"])
    hidden = writer.patch(_student_url(student.id), json={"note": "x"})
    assert hidden.status_code == 200
    assert hidden.json()["sensitive"] is None
    assert hidden.json()["has_health_note"] is True
    for leaked in ("氣喘", _ID_NUMBER, "_enc", "hmac"):
        assert leaked not in hidden.text


def test_admin_students_update_422(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    student = make_student(db_session)
    db_session.commit()
    client, _ = staff_client(permissions=["students:write", "students:sensitive"])

    for body in ({}, {"status": "graduated"}, {"foo": 1}, {"name": None}, {"grade_level": 7}):
        assert_error(client.patch(_student_url(student.id), json=body), 422, "validation_error")
    assert_error(client.patch(_student_url("abc"), json={"note": "x"}), 422, "validation_error")
    assert_error(
        client.patch(_student_url(student.id), json={"id_number": "A123"}), 422, "invalid_id_number"
    )
    assert_error(
        client.patch(
            _student_url(student.id),
            json={"status": "withdrawn", "enrolled_on": "2026-09-01", "withdrawn_on": "2026-08-31"},
        ),
        422,
        "invalid_dates",
    )
    assert _stored(db_session, student.id).status == "active"


def test_admin_students_update_401(api_client: TestClient, assert_error: AssertError) -> None:
    assert_error(
        api_client.patch(_student_url(uuid4()), json={"note": "x"}), 401, "unauthenticated"
    )


def test_admin_students_update_403(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    student = make_student(db_session, name="王小明")
    db_session.commit()
    reader, _ = staff_client(permissions=["students:read"])
    writer, _ = staff_client(permissions=["students:write"])

    denied = reader.patch(_student_url(student.id), json={"note": "x"})
    sensitive = writer.patch(_student_url(student.id), json={"health_note": "氣喘"})
    cleared = writer.patch(_student_url(student.id), json={"id_number": None})

    assert_error(denied, 403, "permission_denied")
    assert denied.json()["error"]["details"] == {"required": ["students:write"]}
    assert_error(sensitive, 403, "sensitive_permission_required")
    assert_error(cleared, 403, "sensitive_permission_required")
    stored = _stored(db_session, student.id)
    assert (stored.note, stored.health_note_enc) == (None, None)


def test_admin_students_update_404_409(
    staff_client: StaffClientFactory, assert_error: AssertError, db_session: Session
) -> None:
    archived = make_student(db_session, name="王小明", archived=True)
    make_student(db_session, student_no="S115001")
    student = make_student(db_session, student_no="S115002")
    db_session.commit()
    client, _ = staff_client(permissions=["students:write"])

    assert_error(
        client.patch(_student_url(archived.id), json={"note": "x"}), 404, "student_not_found"
    )
    assert_error(client.patch(_student_url(uuid4()), json={"note": "x"}), 404, "student_not_found")
    assert_error(
        client.patch(_student_url(student.id), json={"student_no": "S115001"}),
        409,
        "student_no_taken",
    )
    assert _stored(db_session, archived.id).note is None
    assert _stored(db_session, student.id).student_no == "S115002"


def test_admin_students_update_guard_registered(app: FastAPI) -> None:
    assert admin_routes_without_permission(app) == []
    assert "patch" in app.openapi()["paths"]["/api/admin/students/{student_id}"]


# --- BACKEND-533：GET /api/admin/students/import-template -------------------------

_TEMPLATE_URL = f"{_URL}/import-template"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_admin_students_import_template_success(staff_client: StaffClientFactory) -> None:
    client, _ = staff_client(permissions=["students:write"])

    resp = client.get(_TEMPLATE_URL)

    assert resp.status_code == 200
    assert resp.headers["content-type"] == _XLSX
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith("attachment; filename*=UTF-8''")
    assert unquote(disposition.split("''", 1)[1]) == "學生匯入範本.xlsx"
    assert resp.headers["cache-control"] == "no-store"
    wb = load_workbook(BytesIO(resp.content))
    visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    assert visible[0].title == "學生資料"
    assert [c.value for c in visible[0][1]] == IMPORT_COLUMNS
    assert visible[0].max_row == 1
    assert "說明" in wb.sheetnames


def test_admin_students_import_template_route_order(staff_client: StaffClientFactory) -> None:
    client, _ = staff_client(permissions=["students:write"])

    resp = client.get(_TEMPLATE_URL)

    # 沒被 /{student_id} 的 UUID 驗證攔截成 422
    assert resp.status_code == 200
    assert resp.headers["content-type"] == _XLSX


def test_admin_students_import_template_401(
    api_client: TestClient, assert_error: AssertError
) -> None:
    assert_error(api_client.get(_TEMPLATE_URL), 401, "unauthenticated")


def test_admin_students_import_template_403(
    staff_client: StaffClientFactory, assert_error: AssertError
) -> None:
    client, _ = staff_client(permissions=["students:read"])

    resp = client.get(_TEMPLATE_URL)

    assert_error(resp, 403, "permission_denied")
    assert resp.json()["error"]["details"] == {"required": ["students:write"]}


def test_admin_students_import_template_method(
    staff_client: StaffClientFactory, assert_error: AssertError
) -> None:
    client, _ = staff_client(permissions=["students:write"])

    assert_error(client.post(_TEMPLATE_URL), 405, "method_not_allowed")


def test_admin_students_import_template_guard_registered(app: FastAPI) -> None:
    assert admin_routes_without_permission(app) == []
    assert "get" in app.openapi()["paths"]["/api/admin/students/import-template"]
