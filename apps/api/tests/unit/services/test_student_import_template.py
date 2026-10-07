"""BACKEND-532：app/services/student_import_template.py（學生 Excel 匯入範本）。

純函式，不碰 DB：範本以 openpyxl 讀回後逐項斷言標題列、樣式、資料驗證與說明工作表，並以
BACKEND-155 的 check_header / preview / assert_xlsx_within_limits 與 BACKEND-016 的上傳驗證確認
「範本原樣上傳」會被接受為合法檔案、被判定為沒有資料列。
"""

import io
import unicodedata
from typing import Any, cast
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from starlette.datastructures import UploadFile

from app.api.deps import CurrentStaff
from app.core.errors import AppError
from app.core.uploads import XLSX_MAX_BYTES, XLSX_TYPES, ValidatedUpload, read_validated_upload
from app.schemas.students import StudentCreateIn
from app.services.student_import_service import (
    _GENDERS,
    _STATUSES,
    IMPORT_COLUMNS,
    MAX_IMPORT_ROWS,
    assert_xlsx_within_limits,
    check_header,
    preview,
)
from app.services.student_import_template import build_import_template

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook() -> Any:
    return load_workbook(io.BytesIO(build_import_template()))


def _data_sheet() -> Any:
    return _workbook()["學生資料"]


def _letter(column: str) -> str:
    return str(get_column_letter(IMPORT_COLUMNS.index(column) + 1))


def _validations(ws: Any) -> dict[str, Any]:
    """以套用範圍（例如 'C2:C501'）為 key。"""
    return {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}


def _display_width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in text)


def test_import_template_header() -> None:
    ws = _data_sheet()

    assert [cell.value for cell in ws[1]] == IMPORT_COLUMNS
    assert ws["A1"].font.bold is True
    assert ws.freeze_panes == "A2"
    assert ws.max_row == 1


def test_import_template_sheet_order() -> None:
    """BACKEND-155 讀第一張可見工作表：學生資料必須在最前面，說明在後。"""
    wb = _workbook()

    assert wb.sheetnames == ["學生資料", "說明"]
    assert [ws.sheet_state for ws in wb.worksheets] == ["visible", "visible"]


def test_import_template_required_headers_highlighted() -> None:
    ws = _data_sheet()

    for index, column in enumerate(IMPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=index)
        if column.endswith("*"):
            assert cell.font.bold is True, column
            assert cell.fill.fill_type == "solid", column
            assert cell.fill.fgColor.rgb[-6:] == "FFF2CC", column
        else:
            assert not cell.font.bold, column
            assert cell.fill.fill_type is None, column


def test_import_template_column_widths_follow_titles() -> None:
    ws = _data_sheet()

    for index, column in enumerate(IMPORT_COLUMNS, start=1):
        width = ws.column_dimensions[get_column_letter(index)].width
        assert width >= _display_width(column), column
    # 標題較長的欄比較寬
    assert (
        ws.column_dimensions[_letter("安親班班級")].width
        > ws.column_dimensions[_letter("性別")].width
    )


def test_import_template_validations() -> None:
    ws = _data_sheet()
    last_row = MAX_IMPORT_ROWS + 1  # 標題列之後最多可匯入的列數

    validations = _validations(ws)

    gender_col, status_col, grade_col = _letter("性別"), _letter("狀態"), _letter("年級*")
    assert set(validations) == {
        f"{gender_col}2:{gender_col}{last_row}",
        f"{status_col}2:{status_col}{last_row}",
        f"{grade_col}2:{grade_col}{last_row}",
    }
    gender = validations[f"{gender_col}2:{gender_col}{last_row}"]
    assert (gender.type, gender.formula1) == ("list", '"男,女,其他"')
    status = validations[f"{status_col}2:{status_col}{last_row}"]
    assert (status.type, status.formula1) == ("list", '"在學,暫停,退班"')
    grade = validations[f"{grade_col}2:{grade_col}{last_row}"]
    assert (grade.type, grade.operator, grade.formula1, grade.formula2) == (
        "whole",
        "between",
        "1",
        "6",
    )
    for dv in validations.values():
        # 輸入不合法的值時要擋下並說明；空白不在驗證範圍（必填由匯入預覽檢查）
        assert dv.showErrorMessage is True
        assert dv.allow_blank is True
        assert dv.error


def test_import_template_validations_match_importer_rules() -> None:
    """下拉選項與整數範圍必須與匯入實際接受的值一致，匯入規則改動時這裡轉紅。"""
    validations = _validations(_data_sheet())
    last_row = MAX_IMPORT_ROWS + 1

    gender = validations[f"{_letter('性別')}2:{_letter('性別')}{last_row}"]
    status = validations[f"{_letter('狀態')}2:{_letter('狀態')}{last_row}"]
    assert gender.formula1.strip('"').split(",") == list(_GENDERS)
    assert status.formula1.strip('"').split(",") == list(_STATUSES)
    base: dict[str, Any] = {"student_no": "S-001", "name": "王小明"}
    for grade in (1, 6):
        assert StudentCreateIn(**base, grade_level=grade).grade_level == grade
    for grade in (0, 7):
        with pytest.raises(ValueError, match="grade_level"):
            StudentCreateIn(**base, grade_level=grade)


def test_import_template_help_sheet() -> None:
    help_ws = _workbook()["說明"]

    names = [row[0] for row in help_ws.iter_rows(min_row=2, max_col=1, values_only=True)]

    assert set(names) == set(IMPORT_COLUMNS)
    assert len(names) == len(IMPORT_COLUMNS)  # 每欄一列，沒有重複與空列
    assert [cell.value for cell in help_ws[1]][:3] == ["欄名", "是否必填", "格式說明"]


def test_import_template_help_sheet_content() -> None:
    help_ws = _workbook()["說明"]
    rows = {row[0]: row for row in help_ws.iter_rows(min_row=2, max_col=3, values_only=True)}

    for column in IMPORT_COLUMNS:
        assert rows[column][1] == ("必填" if column.endswith("*") else "選填"), column
        assert rows[column][2], column  # 每欄都有格式說明
    assert "YYYY-MM-DD" in rows["生日"][2]
    assert "YYYY-MM-DD" in rows["入學日"][2]
    assert "students:sensitive" in rows["身分證字號"][2]
    assert "students:sensitive" in rows["健康備註"][2]
    assert "名稱" in rows["就讀國小"][2]
    assert "名稱" in rows["安親班班級"][2]
    for option in ("男", "女", "其他"):
        assert option in rows["性別"][2]
    for option in ("在學", "暫停", "退班"):
        assert option in rows["狀態"][2]
    assert "1~6" in rows["年級*"][2]


def test_import_template_passes_header_check() -> None:
    header = [cell.value for cell in _data_sheet()[1]]

    assert check_header(header) == ([], [])


def test_import_template_is_accepted_by_upload_validation() -> None:
    """範本原樣上傳：通過 xlsx 檔頭 / 大小檢查與 zip、XML 層防護。"""
    content = build_import_template()

    upload = read_validated_upload(
        UploadFile(io.BytesIO(content), filename="範本.xlsx"),
        allowed=XLSX_TYPES,
        max_bytes=XLSX_MAX_BYTES,
    )
    assert_xlsx_within_limits(content)

    assert (upload.ext, upload.mime_type, upload.size) == ("xlsx", _XLSX_MIME, len(content))


class _NoDb:
    """範本預覽在標題與資料列階段就結束，不應碰 DB。"""

    def __getattr__(self, name: str) -> Any:
        raise AssertionError(f"範本預覽不應碰 DB：session.{name}")


def test_import_template_preview_reports_empty_not_invalid_header() -> None:
    content = build_import_template()
    actor = CurrentStaff(
        id=uuid4(),
        username="clerk-test",
        display_name="林老師",
        role_id=uuid4(),
        role_code="clerk",
        role_name="行政",
        permissions=frozenset(),
        must_change_password=False,
        token_version=0,
    )
    upload = ValidatedUpload(content=content, mime_type=_XLSX_MIME, ext="xlsx", size=len(content))

    with pytest.raises(AppError) as excinfo:
        preview(cast(Session, _NoDb()), upload, academic_year=115, actor=actor)

    assert (excinfo.value.status, excinfo.value.code) == (422, "import_empty")
