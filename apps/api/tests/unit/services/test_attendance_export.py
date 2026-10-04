"""BACKEND-313：app/services/attendance_export.py（月出勤報表 Excel，純函式）。"""

from __future__ import annotations

import calendar
import datetime as dt
from io import BytesIO
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.attendance import MonthlyAttendanceOut
from app.services.attendance_export import build_monthly_attendance_xlsx

_FIRST_DAY_COLUMN = 3


def _report(
    *,
    name: str = "王小明",
    student_no: str = "S001",
    statuses: dict[int, str] | None = None,
    year: int = 2026,
    month: int = 9,
) -> MonthlyAttendanceOut:
    day_count = calendar.monthrange(year, month)[1]
    days = []
    for n in range(1, day_count + 1):
        d = dt.date(year, month, n)
        days.append({"date": d, "weekday": d.weekday(), "is_service_day": d.weekday() < 5})
    row_statuses: list[str | None] = [None] * day_count
    for idx, status in (statuses or {0: "present", 1: "leave"}).items():
        row_statuses[idx] = status
    stats: dict[str, Any] = {
        "service_days": 22,
        "attended": sum(1 for s in row_statuses if s in ("present", "left")),
        "absent": sum(1 for s in row_statuses if s == "absent"),
        "leave": sum(1 for s in row_statuses if s == "leave"),
        "unrecorded": 5,
    }
    return MonthlyAttendanceOut.model_validate(
        {
            "month": f"{year}-{month:02d}",
            "class_id": uuid4(),
            "class_name": "低年級A班",
            "days": days,
            "students": [
                {
                    "student_id": uuid4(),
                    "student_no": student_no,
                    "name": name,
                    "class_name": "低年級A班",
                    "statuses": row_statuses,
                    "stats": stats,
                }
            ],
            "totals": stats,
        }
    )


def _sheet(report: MonthlyAttendanceOut, title: str = "低年級A班") -> Worksheet:
    wb = load_workbook(BytesIO(build_monthly_attendance_xlsx(report, title_class_name=title)))
    assert wb.sheetnames == ["出勤月報"]
    return wb["出勤月報"]


def test_attendance_xlsx_layout() -> None:
    ws = _sheet(_report())

    assert ws["A1"].value == "低年級A班 2026年9月出勤月報"
    assert "A1:AJ1" in {str(r) for r in ws.merged_cells.ranges}  # 2 + 30 天 + 4 統計 = 36 欄（AJ）
    assert [ws.cell(2, c).value for c in (1, 2, 3, 4, 5)] == [
        "學號",
        "姓名",
        "1日(二)",
        "2日(三)",
        "3日(四)",
    ]
    assert ws["AF2"].value == "30日(三)"
    last = ws.max_column
    assert [ws.cell(2, last - i).value for i in (3, 2, 1, 0)] == [
        "到班",
        "缺席",
        "請假",
        "未登記",
    ]
    assert ws["A3"].value == "S001"
    assert ws["B3"].value == "王小明"
    assert ws["C3"].value == "到"
    assert ws["D3"].value == "假"
    assert ws["E3"].value is None  # None / expected 空白
    assert ws.cell(3, last - 3).value == 1
    assert ws.cell(3, last - 2).value == 0
    assert ws.cell(3, last - 1).value == 1
    assert ws.cell(3, last).value == 5
    assert ws.max_column == 2 + 30 + 4


def test_attendance_xlsx_status_codes() -> None:
    ws = _sheet(_report(statuses={0: "present", 1: "left", 2: "absent", 3: "leave", 6: "expected"}))

    assert ws["C3"].value == "到"
    assert ws["D3"].value == "離"
    assert ws["E3"].value == "缺"
    assert ws["F3"].value == "假"
    assert ws["I3"].value is None  # 9/7 週一 expected → 空白


def test_attendance_xlsx_non_service_day() -> None:
    ws = _sheet(_report())

    sunday = ws.cell(3, _FIRST_DAY_COLUMN + 5)  # 9/6 週日
    service_day = ws.cell(3, _FIRST_DAY_COLUMN)
    assert sunday.value == "-"
    assert sunday.fill.fgColor.rgb.endswith("DDDDDD")
    assert service_day.fill.fgColor.rgb != sunday.fill.fgColor.rgb


def test_attendance_xlsx_formula_injection() -> None:
    ws = _sheet(_report(name='=HYPERLINK("x")', student_no="+1"), title="@班")

    assert ws["B3"].value == '\'=HYPERLINK("x")'
    assert ws["A3"].value == "'+1"
    assert ws["A1"].value == "'@班 2026年9月出勤月報"
    for lead in ("-1", "@a", "\tx"):
        assert _sheet(_report(name=lead)).cell(3, 2).value == "'" + lead
    # xlsx 的 XML 會把 CR 正規化成 LF，只驗證有補前綴
    cr_value = _sheet(_report(name="\rx")).cell(3, 2).value
    assert isinstance(cr_value, str)
    assert cr_value.startswith("'")
    assert cr_value.endswith("x")
    assert _sheet(_report(name="王=小明")).cell(3, 2).value == "王=小明"
    # 非使用者輸入的「-」不可被加前綴
    assert _sheet(_report()).cell(3, _FIRST_DAY_COLUMN + 5).value == "-"


def test_attendance_xlsx_column_width_follows_content() -> None:
    narrow = _sheet(_report(name="王"))
    wide = _sheet(_report(name="王小明" * 8))

    assert wide.column_dimensions["B"].width > narrow.column_dimensions["B"].width
    assert narrow.column_dimensions["B"].width >= 4
