"""BACKEND-155：學生 Excel 匯入預覽 ``preview``（domain_spec §2 ``POST /students/import`` 兩段式的
第一段：解析 + 逐列驗證，不寫入）。移植 ivy ``utils/excel_io.py::parse_excel`` /
``assert_xlsx_not_zip_bomb`` 的防護，欄位改為本專案學生欄位。

- 呼叫端先以 BACKEND-016 ``read_validated_upload``（XLSX 檔頭 + 5 MB 上限）取得
  ``ValidatedUpload``；
  本模組在交給 openpyxl 之前再做 zip 層檢查：解壓後總位元組上限、工作表 XML 的列數 / 單列欄數 /
  總儲存格數上限（避免解壓炸彈與超寬表讓 openpyxl OOM），並丟棄可偽造的 dimension、以硬上限限制
  迭代列數。``data_only=True``：公式儲存格只取快取值，儲存格內容一律當純文字，不評估任何公式。
- 標題列：``check_header`` 回 ``(missing, unexpected)``（必填欄缺少依 IMPORT_COLUMNS 順序、無法辨識
  的欄名依標題列順序；選填欄缺少不算錯）→ 422 ``import_invalid_header``。資料列 > 500 → 422
  ``import_too_many_rows``；沒有資料列（整列空白不算）→ 422 ``import_empty``。
- 逐列：性別 / 狀態中文對照、日期接受 Excel 日期或 ``YYYY-MM-DD`` / ``YYYY/MM/DD``、國小以 name 或
  short_name 比對啟用中的 schools（不分大小寫）、班級以名稱比對該學年度未封存的班、身分證正規化 +
  檢查碼（BACKEND-147）、無 ``students:sensitive`` 不可帶身分證 / 健康備註、學號與身分證 HMAC
  的 DB 重複（含封存）與檔內重複（兩列都標），最後以 ``StudentCreateIn`` 驗證並把 Pydantic 錯誤
  轉成繁中訊息。驗證規則與 BACKEND-151 ``create_student`` 一致（執行階段由 BACKEND-156 共用其
  寫入邏輯）。
- ``display`` 為原始儲存格字串供前端呈現：身分證遮罩（``mask_id_number``）、健康備註以固定文字
  取代。
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Final
from uuid import UUID

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import CurrentStaff
from app.core.errors import AppError
from app.core.permissions import Permission
from app.core.uploads import ValidatedUpload
from app.models.classes import SchoolClass
from app.models.reference import School
from app.models.students import Student
from app.schemas.students import StudentCreateIn
from app.services.students.id_number import (
    id_number_hmac,
    mask_id_number,
    normalize_id_number,
    validate_id_number,
)

IMPORT_COLUMNS: Final[list[str]] = [
    "學號*",
    "姓名*",
    "性別",
    "生日",
    "年級*",
    "就讀國小",
    "學校班級",
    "安親班班級",
    "狀態",
    "入學日",
    "備註",
    "身分證字號",
    "健康備註",
]
MAX_IMPORT_ROWS: Final = 500
MAX_IMPORT_COLS: Final = 64
MAX_XLSX_UNCOMPRESSED_BYTES: Final = 32 * 1024 * 1024
# 標題列 + 少量空白列的寬容值；超過就不再往下讀
_GRID_ROW_SLACK: Final = 16

_REQUIRED_COLUMNS: Final = [c for c in IMPORT_COLUMNS if c.endswith("*")]
_COLUMN_FIELDS: Final[dict[str, str]] = {
    "學號*": "student_no",
    "姓名*": "name",
    "性別": "gender",
    "生日": "birthday",
    "年級*": "grade_level",
    "就讀國小": "school_id",
    "學校班級": "school_class",
    "安親班班級": "class_id",
    "狀態": "status",
    "入學日": "enrolled_on",
    "備註": "note",
    "身分證字號": "id_number",
    "健康備註": "health_note",
}
_FIELD_COLUMNS: Final = {field: column for column, field in _COLUMN_FIELDS.items()}
_GENDERS: Final = {"男": "male", "女": "female", "其他": "other"}
_STATUSES: Final = {"在學": "active", "暫停": "suspended", "退班": "withdrawn"}
_DATE_RE: Final = re.compile(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})$")
_HEALTH_NOTE_MASK: Final = "（已隱藏）"
_SENSITIVE_PERMISSION_ERROR: Final = "沒有權限匯入敏感欄位"

_ROW_TAG: Final = b"<row"
_CELL_TAGS: Final = (b"<c ", b"<c>", b"<c/>")


@dataclass(frozen=True)
class ImportRowResult:
    row_number: int  # Excel 列號（標題為第 1 列）
    data: StudentCreateIn | None  # 驗證通過時的正規化資料
    display: dict[str, str]  # 原始儲存格字串（敏感欄位遮罩），供前端呈現
    errors: list[str]  # 繁中錯誤訊息


@dataclass(frozen=True)
class ImportPreview:
    rows: list[ImportRowResult]
    total: int
    valid: int
    invalid: int


def _invalid_file(message: str) -> AppError:
    return AppError("import_invalid_file", message, status=422)


# --- 檔案層防護 ------------------------------------------------------------------------


def _assert_sheet_grid_within_limits(data: bytes) -> None:
    """掃描工作表 XML：列數 / 總儲存格數 / 單列欄數皆在上限內（bytes.count，不切片複製）。"""
    n_rows = data.count(_ROW_TAG)
    if n_rows > MAX_IMPORT_ROWS + _GRID_ROW_SLACK:
        raise AppError(
            "import_too_many_rows",
            f"匯入列數超過上限 {MAX_IMPORT_ROWS}，請分批匯入",
            status=422,
            details={"max_rows": MAX_IMPORT_ROWS},
        )
    if sum(data.count(tag) for tag in _CELL_TAGS) > MAX_IMPORT_ROWS * MAX_IMPORT_COLS:
        raise _invalid_file("Excel 儲存格數量超過上限，請確認檔案內容")
    starts: list[int] = []
    cursor = 0
    while (idx := data.find(_ROW_TAG, cursor)) >= 0:
        starts.append(idx)
        cursor = idx + len(_ROW_TAG)
    for i, start in enumerate(starts):
        end = starts[i + 1] if i + 1 < len(starts) else len(data)
        if sum(data.count(tag, start, end) for tag in _CELL_TAGS) > MAX_IMPORT_COLS:
            raise _invalid_file(f"Excel 欄位數超過上限 {MAX_IMPORT_COLS}，請確認檔案未含異常寬列")


def assert_xlsx_within_limits(content: bytes) -> None:
    """交給 openpyxl 之前：解壓後總大小與工作表格數上限（解壓炸彈 / 超寬表）。"""
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
    except (zipfile.BadZipFile, OSError, ValueError):
        raise _invalid_file("無法讀取 Excel 檔案") from None
    with zf:
        total = 0
        for info in zf.infolist():
            total += info.file_size
            if total > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise _invalid_file("Excel 解壓後大小超過上限（疑似壓縮炸彈），請確認檔案內容")
        for info in zf.infolist():
            name = info.filename.replace("\\", "/").lower()
            if "worksheets/" in name and name.endswith(".xml"):
                _assert_sheet_grid_within_limits(zf.read(info))


# --- 標題與讀列 ------------------------------------------------------------------------


def _header_text(cell: object) -> str:
    return "" if cell is None else str(cell).strip()


def check_header(header_cells: Sequence[str | None]) -> tuple[list[str], list[str]]:
    """回 ``(missing, unexpected)``：缺少的必填欄（IMPORT_COLUMNS 順序）、無法辨識的非空白欄名
    （標題列順序）。選填欄缺少不算錯。"""
    present = [_header_text(c) for c in header_cells]
    missing = [c for c in _REQUIRED_COLUMNS if c not in present]
    unexpected = [c for c in present if c and c not in IMPORT_COLUMNS]
    return missing, unexpected


def _read_rows(content: bytes) -> tuple[list[str], list[tuple[int, tuple[Any, ...]]]]:
    """第一個工作表 → (標題, [(Excel 列號, 儲存格值...)])；整列空白略過。"""
    assert_xlsx_within_limits(content)
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # openpyxl 對壞檔會拋各種解析例外，一律視為無法讀取
        raise _invalid_file("無法讀取 Excel 檔案") from None
    try:
        ws = wb.worksheets[0] if wb.worksheets else None
        if ws is None:
            raise AppError("import_empty", "Excel 沒有資料列", status=422)
        # 丟棄可偽造的 dimension（否則 read_only 會依宣告的末列補出大量空白列），並以硬上限限制
        # 迭代量
        ws.reset_dimensions()
        rows_iter = ws.iter_rows(values_only=True, max_row=MAX_IMPORT_ROWS + _GRID_ROW_SLACK + 1)
        header_row = next(rows_iter, None)
        if header_row is None or not any(_header_text(c) for c in header_row):
            raise AppError("import_empty", "Excel 沒有資料列", status=422)
        header = [_header_text(c) for c in header_row]
        if len(header) > MAX_IMPORT_COLS:
            raise _invalid_file(f"Excel 欄位數超過上限 {MAX_IMPORT_COLS}")
        missing, unexpected = check_header(header)
        if missing or unexpected:
            raise AppError(
                "import_invalid_header",
                "Excel 標題列與匯入範本不符",
                status=422,
                details={"missing": missing, "unexpected": unexpected},
            )
        rows: list[tuple[int, tuple[Any, ...]]] = []
        for row_number, raw in enumerate(rows_iter, start=2):
            if raw is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in raw):
                continue
            rows.append((row_number, tuple(raw)))
            if len(rows) > MAX_IMPORT_ROWS:
                raise AppError(
                    "import_too_many_rows",
                    f"匯入列數超過上限 {MAX_IMPORT_ROWS}，請分批匯入",
                    status=422,
                    details={"max_rows": MAX_IMPORT_ROWS},
                )
    finally:
        wb.close()
    if not rows:
        raise AppError("import_empty", "Excel 沒有資料列", status=422)
    return header, rows


# --- 儲存格轉換 ------------------------------------------------------------------------


def _cell_text(value: object) -> str | None:
    """儲存格 → 去頭尾空白的字串；空白視為未填。整數值的浮點數（Excel 數字）去掉 .0。"""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        value = value.date()
    text = str(value).strip()
    return text or None


def _parse_date(value: object, column: str, errors: list[str]) -> date | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = _DATE_RE.match(str(value).strip())
    if match is not None:
        try:
            return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        except ValueError:
            pass
    errors.append(f"{column}格式不正確（請用 YYYY-MM-DD）")
    return None


def _parse_choice(
    text: str | None, mapping: Mapping[str, str], column: str, errors: list[str]
) -> str | None:
    if text is None:
        return None
    mapped = mapping.get(text)
    if mapped is None:
        errors.append(f"{column}只能填 {' / '.join(mapping)}")
    return mapped


def _parse_grade(text: str | None, errors: list[str]) -> int | None:
    if text is None:
        return None
    if not text.isdigit():
        errors.append("年級*必須是 1~6 的整數")
        return None
    return int(text)


def _display_value(column: str, value: object) -> str:
    text = _cell_text(value)
    if text is None:
        return ""
    if column == "身分證字號":
        normalized = normalize_id_number(text)
        return mask_id_number(normalized) if len(normalized) >= 7 else "***"
    if column == "健康備註":
        return _HEALTH_NOTE_MASK
    return text


def _validation_messages(exc: ValidationError) -> list[str]:
    messages: list[str] = []
    for err in exc.errors():
        loc = err.get("loc", ())
        field = str(loc[0]) if loc else ""
        column = _FIELD_COLUMNS.get(field, field)
        messages.append(f"{column}格式不正確：{err.get('msg', '')}")
    return messages


# --- 查表 ---------------------------------------------------------------------------


@dataclass(frozen=True)
class _Lookups:
    schools: dict[str, UUID]  # casefold(name / short_name) → id（只含啟用中）
    classes: dict[str, UUID]  # casefold(name) → id（該學年度未封存）
    existing_student_nos: set[str]
    existing_hmacs: set[str]


def _load_lookups(
    session: Session, *, academic_year: int, student_nos: Iterable[str], hmacs: Iterable[str]
) -> _Lookups:
    schools: dict[str, UUID] = {}
    for school in session.execute(select(School).where(School.is_active.is_(True))).scalars():
        schools[school.name.strip().casefold()] = school.id
        if school.short_name:
            schools.setdefault(school.short_name.strip().casefold(), school.id)
    classes = {
        name.strip().casefold(): class_id
        for class_id, name in session.execute(
            select(SchoolClass.id, SchoolClass.name).where(
                SchoolClass.academic_year == academic_year, SchoolClass.archived_at.is_(None)
            )
        )
    }
    nos = list(set(student_nos))
    existing_nos: set[str] = set()
    if nos:
        existing_nos = set(
            session.execute(select(Student.student_no).where(Student.student_no.in_(nos))).scalars()
        )
    hmac_list = list(set(hmacs))
    existing_hmacs: set[str] = set()
    if hmac_list:
        existing_hmacs = {
            h
            for h in session.execute(
                select(Student.id_number_hmac).where(Student.id_number_hmac.in_(hmac_list))
            ).scalars()
            if h is not None
        }
    return _Lookups(schools, classes, existing_nos, existing_hmacs)


# --- preview ----------------------------------------------------------------------


def _cells(header: list[str], raw: tuple[Any, ...]) -> dict[str, Any]:
    """依標題名取值（忽略空白標題欄）。"""
    return {name: (raw[i] if i < len(raw) else None) for i, name in enumerate(header) if name}


def preview(
    session: Session, upload: ValidatedUpload, *, academic_year: int, actor: CurrentStaff
) -> ImportPreview:
    header, raw_rows = _read_rows(upload.content)
    rows = [(row_number, _cells(header, raw)) for row_number, raw in raw_rows]
    can_sensitive = actor.has(Permission.STUDENTS_SENSITIVE)

    # 先算學號與身分證 HMAC：一次查 DB、同時找出檔內重複
    student_nos = [_cell_text(cells.get("學號*")) for _, cells in rows]
    normalized_ids: list[str | None] = []
    hmacs: list[str | None] = []
    for _, cells in rows:
        raw_id = _cell_text(cells.get("身分證字號"))
        normalized = normalize_id_number(raw_id) if raw_id is not None else None
        normalized_ids.append(normalized)
        hmacs.append(id_number_hmac(normalized) if normalized is not None else None)
    lookups = _load_lookups(
        session,
        academic_year=academic_year,
        student_nos=[n for n in student_nos if n is not None],
        hmacs=[h for h in hmacs if h is not None],
    )
    no_counts: dict[str, int] = {}
    for no in student_nos:
        if no is not None:
            no_counts[no] = no_counts.get(no, 0) + 1
    hmac_counts: dict[str, int] = {}
    for hmac in hmacs:
        if hmac is not None:
            hmac_counts[hmac] = hmac_counts.get(hmac, 0) + 1

    results: list[ImportRowResult] = []
    for index, (row_number, cells) in enumerate(rows):
        errors: list[str] = []
        display = {column: _display_value(column, cells.get(column)) for column in IMPORT_COLUMNS}

        student_no = student_nos[index]
        if student_no is not None:
            if student_no in lookups.existing_student_nos:
                errors.append("學號已存在")
            if no_counts.get(student_no, 0) > 1:
                errors.append("學號在檔案中重複")

        school_text = _cell_text(cells.get("就讀國小"))
        school_id: UUID | None = None
        if school_text is not None:
            school_id = lookups.schools.get(school_text.casefold())
            if school_id is None:
                errors.append(f"找不到國小：{school_text}")
        class_text = _cell_text(cells.get("安親班班級"))
        class_id: UUID | None = None
        if class_text is not None:
            class_id = lookups.classes.get(class_text.casefold())
            if class_id is None:
                errors.append(f"找不到安親班班級：{class_text}（{academic_year} 學年度）")

        health_note = _cell_text(cells.get("健康備註"))
        normalized_id = normalized_ids[index]
        if (normalized_id is not None or health_note is not None) and not can_sensitive:
            errors.append(_SENSITIVE_PERMISSION_ERROR)
        if normalized_id is not None:
            try:
                validate_id_number(normalized_id)
            except AppError:
                errors.append("身分證字號格式不正確")
            else:
                hmac = hmacs[index]
                if hmac in lookups.existing_hmacs:
                    errors.append("身分證字號已存在")
                if hmac is not None and hmac_counts.get(hmac, 0) > 1:
                    errors.append("身分證字號在檔案中重複")

        payload: dict[str, Any] = {
            "student_no": student_no,
            "name": _cell_text(cells.get("姓名*")),
            "gender": _parse_choice(_cell_text(cells.get("性別")), _GENDERS, "性別", errors),
            "birthday": _parse_date(cells.get("生日"), "生日", errors),
            "grade_level": _parse_grade(_cell_text(cells.get("年級*")), errors),
            "school_id": school_id,
            "school_class": _cell_text(cells.get("學校班級")),
            "class_id": class_id,
            "status": _parse_choice(_cell_text(cells.get("狀態")), _STATUSES, "狀態", errors)
            or "active",
            "enrolled_on": _parse_date(cells.get("入學日"), "入學日", errors),
            "note": _cell_text(cells.get("備註")),
            "id_number": normalized_id,
            "health_note": health_note,
        }
        data: StudentCreateIn | None = None
        try:
            data = StudentCreateIn.model_validate(payload)
        except ValidationError as exc:
            errors.extend(_validation_messages(exc))
        if errors:
            data = None
        results.append(
            ImportRowResult(row_number=row_number, data=data, display=display, errors=errors)
        )

    valid = sum(1 for r in results if r.data is not None)
    return ImportPreview(
        rows=results, total=len(results), valid=valid, invalid=len(results) - valid
    )


__all__ = [
    "IMPORT_COLUMNS",
    "MAX_IMPORT_ROWS",
    "ImportPreview",
    "ImportRowResult",
    "assert_xlsx_within_limits",
    "check_header",
    "preview",
]
