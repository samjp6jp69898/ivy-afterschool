"""BACKEND-313：月出勤報表 Excel 匯出（純函式，不碰 DB）。

版面移植 ivy ``api/student_attendance.py::_write_class_sheet``（單一班級工作表，去掉全園摘要與
中文狀態欄）。使用者輸入字串（班名、學號、姓名）寫入前先做公式注入防護：以 ``=``、``+``、``-``、
``@``、Tab、CR 開頭者前面補 ``'``。程式自己產生的常數（表頭、簡碼、「-」）不經防護。
"""

from __future__ import annotations

import unicodedata
from io import BytesIO
from typing import Final

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.attendance import MonthlyAttendanceOut

SHEET_NAME: Final = "出勤月報"
_FORMULA_PREFIXES: Final = ("=", "+", "-", "@", "\t", "\r")
_WEEKDAYS: Final = "一二三四五六日"
_STATUS_CODES: Final = {"present": "到", "left": "離", "absent": "缺", "leave": "假"}
_NON_SERVICE_FILL: Final = PatternFill("solid", fgColor="DDDDDD")
_STAT_HEADERS: Final = ("到班", "缺席", "請假", "未登記")
_MIN_WIDTH: Final = 4
_MAX_WIDTH: Final = 40


def _safe(value: str) -> str:
    return "'" + value if value.startswith(_FORMULA_PREFIXES) else value


def _display_width(value: object) -> int:
    text = "" if value is None else str(value)
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def build_monthly_attendance_xlsx(report: MonthlyAttendanceOut, *, title_class_name: str) -> bytes:
    year, month = (int(part) for part in report.month.split("-"))
    wb = Workbook()
    ws = wb.active
    assert isinstance(ws, Worksheet)  # noqa: S101  新 Workbook 的 active 必為 worksheet
    ws.title = SHEET_NAME

    day_count = len(report.days)
    last_col = 2 + day_count + len(_STAT_HEADERS)

    ws.cell(1, 1, _safe(f"{title_class_name} {year}年{month}月出勤月報"))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["學號", "姓名"]
    headers += [f"{d.date.day}日({_WEEKDAYS[d.weekday]})" for d in report.days]
    headers += list(_STAT_HEADERS)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(2, col, header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_idx, student in enumerate(report.students, start=3):
        ws.cell(row_idx, 1, _safe(student.student_no))
        ws.cell(row_idx, 2, _safe(student.name))
        for offset, (day, status) in enumerate(zip(report.days, student.statuses, strict=True)):
            cell = ws.cell(row_idx, 3 + offset)
            if not day.is_service_day:
                cell.value = "-"
                cell.fill = _NON_SERVICE_FILL
            else:
                cell.value = _STATUS_CODES.get(status) if status else None
            cell.alignment = Alignment(horizontal="center")
        stats = student.stats
        for offset, value in enumerate(
            (stats.attended, stats.absent, stats.leave, stats.unrecorded)
        ):
            ws.cell(row_idx, 3 + day_count + offset, value)

    _fit_columns(ws, last_col)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _fit_columns(ws: Worksheet, last_col: int) -> None:
    # 第 1 列是跨欄合併的標題，不參與欄寬計算
    for col in range(1, last_col + 1):
        cells: list[Cell] = [ws.cell(row, col) for row in range(2, ws.max_row + 1)]
        widest = max((_display_width(c.value) for c in cells), default=0)
        ws.column_dimensions[get_column_letter(col)].width = min(
            max(widest + 2, _MIN_WIDTH), _MAX_WIDTH
        )
