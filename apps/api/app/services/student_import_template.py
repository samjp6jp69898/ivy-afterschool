"""BACKEND-532：學生 Excel 匯入範本 ``build_import_template``（domain_spec M3：Excel 匯入只支援後台
提供的範本，範本由後端依匯入欄位定義產生）。純函式，不碰 DB。

- 工作表「學生資料」（必須排第一張：BACKEND-155 讀第一張可見工作表）：第 1 列是 ``IMPORT_COLUMNS``
  （與匯入共用同一個常數，標題列永遠通過 ``check_header``），必填欄（帶 ``*``）粗體 + 淡黃底；
  凍結標題列；欄寬依標題；沒有資料列，所以範本原樣上傳會得到 422 ``import_empty``（標題通過
  檢查），而不是 ``import_invalid_header``。
- 資料驗證只套用在可匯入的列（第 2 列起共 ``MAX_IMPORT_ROWS`` 列）：性別、狀態為下拉選單，年級為
  1~6 的整數。選項與範圍必須與匯入實際接受的值一致（由測試對照 ``student_import_service`` 的對照表與
  ``StudentCreateIn``）。
- 工作表「說明」：逐欄列出欄名、是否必填與格式說明。
"""

from __future__ import annotations

import io
import unicodedata
from typing import Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.services.student_import_service import IMPORT_COLUMNS, MAX_IMPORT_ROWS

DATA_SHEET: Final = "學生資料"
HELP_SHEET: Final = "說明"

GENDER_OPTIONS: Final = ("男", "女", "其他")
STATUS_OPTIONS: Final = ("在學", "暫停", "退班")
GRADE_MIN: Final = 1
GRADE_MAX: Final = 6

_GENDER_COLUMN: Final = "性別"
_STATUS_COLUMN: Final = "狀態"
_GRADE_COLUMN: Final = "年級*"

_REQUIRED_MARK: Final = "*"
_REQUIRED_FILL: Final = PatternFill(fill_type="solid", start_color="FFF2CC", end_color="FFF2CC")
_BOLD: Final = Font(bold=True)
_MIN_COLUMN_WIDTH: Final = 12
_WIDTH_PADDING: Final = 4

_HELP_HEADER: Final = ("欄名", "是否必填", "格式說明")
_HELP: Final[dict[str, str]] = {
    "學號*": "全校唯一，不可與既有學生重複；只能用英文字母、數字與 -，最多 20 碼。",
    "姓名*": "最多 50 字。",
    "性別": f"{' / '.join(GENDER_OPTIONS)}（可用下拉選單選擇），可留空。",
    "生日": "日期，格式 YYYY-MM-DD（例如 2016-09-01），也可直接填 Excel 日期，可留空。",
    "年級*": f"{GRADE_MIN}~{GRADE_MAX} 的整數（國小年級）。",
    "就讀國小": "需與系統內合作國小的名稱或簡稱一致（只比對啟用中的國小），可留空。",
    "學校班級": "學生在國小的班級，例如「三年二班」，最多 20 字，可留空。",
    "安親班班級": "需與系統內該學年度安親班班級的名稱一致（已封存的班級不可用），可留空。",
    "狀態": f"{' / '.join(STATUS_OPTIONS)}（可用下拉選單選擇），留空視為在學。",
    "入學日": "日期，格式 YYYY-MM-DD，也可直接填 Excel 日期，可留空。",
    "備註": "最多 500 字，可留空。",
    "身分證字號": "需有 students:sensitive 權限才能匯入；格式與檢查碼需正確，不可與既有學生重複。",
    "健康備註": "需有 students:sensitive 權限才能匯入；最多 1000 字，可留空。",
}


def _display_width(text: str) -> int:
    """全形（中日韓）字元佔 2 個半形寬。"""
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in text)


def _column_width(title: str) -> int:
    return max(_MIN_COLUMN_WIDTH, _display_width(title) + _WIDTH_PADDING)


def _letter(column: str) -> str:
    # 欄名不在 IMPORT_COLUMNS 時 index 會拋 ValueError：匯入改欄名而這裡沒跟上時直接失敗
    return str(get_column_letter(IMPORT_COLUMNS.index(column) + 1))


def _importable_range(column: str) -> str:
    letter = _letter(column)
    return f"{letter}2:{letter}{MAX_IMPORT_ROWS + 1}"


def _list_validation(column: str, options: tuple[str, ...]) -> DataValidation:
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(options) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle=f"{column}不正確",
        error=f"{column}只能填 {' / '.join(options)}",
    )
    validation.add(_importable_range(column))
    return validation


def _grade_validation() -> DataValidation:
    validation = DataValidation(
        type="whole",
        operator="between",
        formula1=str(GRADE_MIN),
        formula2=str(GRADE_MAX),
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="年級不正確",
        error=f"年級必須是 {GRADE_MIN}~{GRADE_MAX} 的整數",
    )
    validation.add(_importable_range(_GRADE_COLUMN))
    return validation


def _fill_data_sheet(ws: Worksheet) -> None:
    ws.title = DATA_SHEET
    for index, column in enumerate(IMPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=index, value=column)
        if column.endswith(_REQUIRED_MARK):
            cell.font = _BOLD
            cell.fill = _REQUIRED_FILL
        ws.column_dimensions[get_column_letter(index)].width = _column_width(column)
    ws.freeze_panes = "A2"
    ws.add_data_validation(_list_validation(_GENDER_COLUMN, GENDER_OPTIONS))
    ws.add_data_validation(_list_validation(_STATUS_COLUMN, STATUS_OPTIONS))
    ws.add_data_validation(_grade_validation())


def _fill_help_sheet(ws: Worksheet) -> None:
    ws.title = HELP_SHEET
    for index, title in enumerate(_HELP_HEADER, start=1):
        ws.cell(row=1, column=index, value=title).font = _BOLD
    for row, column in enumerate(IMPORT_COLUMNS, start=2):
        required = column.endswith(_REQUIRED_MARK)
        ws.cell(row=row, column=1, value=column)
        ws.cell(row=row, column=2, value="必填" if required else "選填")
        ws.cell(row=row, column=3, value=_HELP[column]).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = max(_column_width(c) for c in IMPORT_COLUMNS)
    ws.column_dimensions["B"].width = _MIN_COLUMN_WIDTH
    ws.column_dimensions["C"].width = 70


def build_import_template() -> bytes:
    workbook = Workbook()
    _fill_data_sheet(workbook.active)
    _fill_help_sheet(workbook.create_sheet())
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
